"""Serialização compartilhada dos downloads CSV/XLSX.

Convenções:
- ausência é célula vazia no arquivo, nunca o rótulo cinza da tela.
- múltiplos nºs SEI concatenados numa única célula, separados por espaço.
- CSV: separador `;`, BOM UTF-8, datas DD/MM/AAAA, decimal com vírgula.
- XLSX: Decimal/int/date nativos com number_format, autofiltro, cabeçalho
  congelado e linha de total com SUM() real sobre colunas monetárias.
- todo campo textual passa pela neutralização de formula injection.
"""

import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from apps.pca.colunas import COLUNAS as _COLUNAS_TELA
from apps.pca.colunas import canonizar
from apps.pca.models import Estado, Situacao

# fallback usado só quando exercicio=None é passado
NOME_CSV = "pca-2026-processos.csv"
NOME_XLSX = "pca-2026-processos.xlsx"

# sem prefixo "R$"; number_format nativo do openpyxl sobre células Decimal
FORMATO_MOEDA = "#,##0.00"
FORMATO_DATA = "dd/mm/yyyy"

# prefixos que o Excel interpreta como início de fórmula
PREFIXOS_FORMULA = ("=", "+", "-", "@")

TIPO_TEXTO = "texto"
TIPO_INTEIRO = "inteiro"
TIPO_DECIMAL = "decimal"
TIPO_DATA = "data"

# vocabulário de tipo da tela mapeado ao vocabulário interno deste módulo;
# status/booleano viram texto porque _extrair já devolve a string formatada
_MAPA_TIPO_TELA = {
    "numero": TIPO_INTEIRO,
    "moeda": TIPO_DECIMAL,
    "data": TIPO_DATA,
    "texto": TIPO_TEXTO,
    "status": TIPO_TEXTO,
    "booleano": TIPO_TEXTO,
}

# mes_previsto é "texto" na tela (filtro de nome de mês), mas o dado em si
# é o inteiro 1-12 — precisa ordenar/somar corretamente no Excel
_EXCECOES_TIPO = {"mes_previsto": TIPO_INTEIRO}


def _cabecalho_export(coluna):
    """Rótulo do arquivo = rótulo da tela, com a unidade monetária explícita
    nas colunas de valor."""
    if coluna.tipo == "moeda":
        return f"{coluna.rotulo} (R$)"
    return coluna.rotulo


# fonte única de colunas: apps.pca.colunas.COLUNAS, mapeada ao vocabulário
# de tipo interno; a tupla preserva (chave, cabeçalho, tipo)
COLUNAS = tuple(
    (
        coluna.chave,
        _cabecalho_export(coluna),
        _EXCECOES_TIPO.get(coluna.chave, _MAPA_TIPO_TELA[coluna.tipo]),
    )
    for coluna in _COLUNAS_TELA
)

# allowlist do export por chave, derivada de COLUNAS
COLUNAS_POR_CHAVE = {chave: (chave, cabecalho, tipo) for chave, cabecalho, tipo in COLUNAS}

# larguras (unidades Excel) por chave, lidas do mesmo registro da tela
_LARGURAS_XLSX = {coluna.chave: coluna.largura_xlsx for coluna in _COLUNAS_TELA}


def _largura_xlsx(chave):
    return _LARGURAS_XLSX.get(chave, 18)


def _colunas_resolvidas(colunas_selecionadas):
    """Resolve o subconjunto de COLUNAS a exportar, na ordem canônica de
    canonizar(). Sem colunas_selecionadas, devolve COLUNAS inteiro; se todas
    as chaves recebidas forem inválidas, cai de volta em COLUNAS inteiro
    para nunca gerar tabela/arquivo sem colunas."""
    if not colunas_selecionadas:
        return COLUNAS
    vistas = set()
    validas = []
    for chave in colunas_selecionadas:
        if chave in COLUNAS_POR_CHAVE and chave not in vistas:
            vistas.add(chave)
            validas.append(chave)
    if not validas:
        return COLUNAS
    return tuple(COLUNAS_POR_CHAVE[chave] for chave in canonizar(validas))


def _texto_seguro(valor):
    """Neutraliza formula injection: prefixa com apóstrofo texto que começa
    com =, +, - ou @."""
    if isinstance(valor, str) and valor.lstrip()[:1] in PREFIXOS_FORMULA:
        return "'" + valor
    return valor


def _extrair(processo, chave):
    """Devolve o valor Python nativo da coluna para o processo, com None
    para toda ausência; sem nenhuma consulta adicional por linha."""
    if chave == "tipo":
        valor = processo.tipo.nome
    elif chave == "categoria":
        valor = processo.categoria.nome
    elif chave == "unidade_organizacional":
        valor = processo.unidade_organizacional.nome
    elif chave == "grau_prioridade":
        valor = processo.grau_prioridade.nome if processo.grau_prioridade else None
    elif chave == "classificacao":
        valor = processo.classificacao.nome if processo.classificacao else None
    elif chave == "modalidade":
        valor = processo.modalidade.nome if processo.modalidade else None
    elif chave == "instrumento_contratual":
        valor = (
            processo.instrumento_contratual.nome
            if processo.instrumento_contratual
            else None
        )
    elif chave == "estado":
        valor = processo.get_estado_display()
    elif chave == "situacao":
        # exporta situacao_efetiva (mesma correção-na-leitura da tela);
        # estado é checado antes, mesma regra da tela
        if processo.estado == Estado.CANCELADO.value:
            valor = "Cancelado"
        else:
            valor = Situacao(processo.situacao_efetiva).label
    elif chave == "numeros_sei":
        # achatado numa célula, separado por um espaço, sem consultas aqui
        valor = " ".join(sei.numero_sei for sei in processo.numeros_sei.all())
    elif chave == "situacao_sei":
        valor = processo.get_situacao_sei_display() if processo.situacao_sei else None
    elif chave in ("atrasado", "publicacao_pendente"):
        # anotações booleanas nunca None: saída sempre "Sim" ou "Não"
        valor = "Sim" if getattr(processo, chave) else "Não"
    else:
        valor = getattr(processo, chave)
    # string vazia também é ausência: no XLSX seria célula de texto, não vazio real
    if valor == "":
        return None
    return valor


# resultado vazio ganha indicação explícita na primeira célula de dados
MENSAGEM_SEM_RESULTADOS = "Nenhum processo encontrado para os filtros aplicados."


def _nome_arquivo(extensao, exercicio):
    """Nome do arquivo pelo exercício efetivo; sem exercicio, cai no literal padrão."""
    if exercicio is not None:
        return f"pca-{exercicio.ano}-processos.{extensao}"
    return NOME_CSV if extensao == "csv" else NOME_XLSX


def _titulo_aba(exercicio):
    """Título da aba pelo exercício efetivo, truncado a 31 caracteres (limite do Excel)."""
    if exercicio is not None:
        return f"Processos PCA {exercicio.ano}"[:31]
    return "Processos PCA 2026"


def exportar_csv(queryset, colunas_selecionadas=None, exercicio=None):
    """CSV `;` com BOM UTF-8, datas DD/MM/AAAA e decimal com vírgula. Uma
    linha por processo, sem linha de total."""
    colunas = _colunas_resolvidas(colunas_selecionadas)
    buffer = StringIO()
    # utf-8-sig: BOM faz o Excel pt-BR escolher UTF-8 em vez de Windows-1252
    buffer.write("\ufeff")
    escritor = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    escritor.writerow([cabecalho for _, cabecalho, _ in colunas])
    for processo in queryset:
        linha = []
        for chave, _, tipo in colunas:
            valor = _extrair(processo, chave)
            if valor is None:
                linha.append("")
            elif tipo == TIPO_DATA:
                linha.append(valor.strftime("%d/%m/%Y"))
            elif tipo == TIPO_DECIMAL:
                # vírgula decimal sem separador de milhar
                linha.append(f"{valor:.2f}".replace(".", ","))
            elif tipo == TIPO_INTEIRO:
                linha.append(str(valor))
            else:
                linha.append(_texto_seguro(str(valor)))
        escritor.writerow(linha)
    resposta = HttpResponse(buffer.getvalue(), content_type="text/csv; charset=utf-8")
    nome_arquivo = _nome_arquivo("csv", exercicio)
    resposta["Content-Disposition"] = f'attachment; filename="{nome_arquivo}"'
    return resposta


def exportar_xlsx(queryset, colunas_selecionadas=None, exercicio=None):
    """XLSX com tipos nativos, number_format monetário/data, larguras de
    coluna, quebra de linha em texto longo, autofiltro, cabeçalho congelado
    e linha de total com SUM() real por coluna monetária."""
    colunas = _colunas_resolvidas(colunas_selecionadas)
    planilha = Workbook()
    ws = planilha.active
    ws.title = _titulo_aba(exercicio)
    ws.append([cabecalho for _, cabecalho, _ in colunas])

    # largura por coluna, antes do loop de linhas
    for indice, (chave, _, _) in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = _largura_xlsx(chave)

    processos = list(queryset)
    for processo in processos:
        ws.append(
            [
                _texto_seguro(_extrair(processo, chave))
                for chave, _, _ in colunas
            ]
        )

    if not processos:
        # resultado vazio: cabeçalho + indicação de ausência, sem linha de total
        ws.append([MENSAGEM_SEM_RESULTADOS])

    ultima_linha_dados = 1 + len(processos)
    ultima_linha = ws.max_row

    # formatos nativos por coluna, aplicados célula a célula; só sobre linhas de dado real
    for indice, (_, _, tipo) in enumerate(colunas):
        letra = get_column_letter(indice + 1)
        if tipo in (TIPO_DECIMAL, TIPO_DATA):
            formato = FORMATO_MOEDA if tipo == TIPO_DECIMAL else FORMATO_DATA
            for linha in range(2, ultima_linha_dados + 1):
                ws[f"{letra}{linha}"].number_format = formato
        elif tipo == TIPO_TEXTO:
            # quebra de linha em texto longo (Objeto/Justificativa etc.)
            for linha in range(2, ultima_linha_dados + 1):
                ws[f"{letra}{linha}"].alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

    # autofiltro e congelamento cobrem cabeçalho + dados; total fica de fora
    ultima_coluna = get_column_letter(len(colunas))
    ws.auto_filter.ref = f"A1:{ultima_coluna}{ultima_linha}"
    ws.freeze_panes = "A2"

    # linha de total: SUM real por coluna monetária, rótulo "Total" na
    # primeira célula; export sem dado nenhum não ganha total
    colunas_decimais = [
        (indice, chave)
        for indice, (chave, _, tipo) in enumerate(colunas, start=1)
        if tipo == TIPO_DECIMAL
    ]
    if processos and colunas_decimais:
        linha_total = ultima_linha_dados + 1
        ws.cell(row=linha_total, column=1, value="Total")
        for indice, _chave in colunas_decimais:
            letra = get_column_letter(indice)
            celula_total = ws.cell(
                row=linha_total,
                column=indice,
                value=f"=SUM({letra}2:{letra}{ultima_linha_dados})",
            )
            celula_total.number_format = FORMATO_MOEDA

    buffer = BytesIO()
    planilha.save(buffer)
    resposta = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    nome_arquivo = _nome_arquivo("xlsx", exercicio)
    resposta["Content-Disposition"] = f'attachment; filename="{nome_arquivo}"'
    return resposta
