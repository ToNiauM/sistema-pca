"""Fluxo administrativo de importação da planilha."""

import os
import tempfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.contrib import admin
from django.core.management.base import CommandError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.http import HttpResponse, HttpResponseRedirect
from django.test import Client, TransactionTestCase, override_settings
from django.urls import path, reverse
from openpyxl import load_workbook

from apps.catalogo.models import Categoria, Exercicio
from apps.pca.importacao.executor import COLUNAS_CATALOGO
from apps.pca.importacao import views
from apps.pca.importacao.models import EventoImportacao
from apps.pca.models import Acompanhamento, HistoricalProcesso, Processo


urlpatterns = [
    path("admin/importar/", views.importar_view, name="importar-teste"),
    path("admin/importar/confirmar/", views.confirmar_view, name="confirmar-teste"),
    path("admin/", admin.site.urls),
]

ARQUIVO_REAL = Path("apps/pca/fixtures/modelo-controle-exemplo.xlsx")


@override_settings(ROOT_URLCONF="apps.pca.tests.test_import_admin")
class TestImportacaoAdmin(TransactionTestCase):
    """Views sem depender das URLs do ModelAdmin, cobertas na Task 2."""

    serialized_rollback = True

    def setUp(self):
        self.client = Client()
        self.superuser = get_user_model().objects.create_superuser(
            email="admin@pca.local", password="senha-segura"
        )
        self.staff = get_user_model().objects.create_user(
            email="staff@pca.local", password="senha-segura", is_staff=True
        )
        self.exercicio = Exercicio.objects.get(ano=2026)
        self.conteudo = ARQUIVO_REAL.read_bytes()

    def _arquivo(self, nome="modelo-controle-exemplo.xlsx", conteudo=None):
        return SimpleUploadedFile(
            nome, self.conteudo if conteudo is None else conteudo,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @staticmethod
    def _renderizar(_request, _template, contexto):
        relatorio = contexto.get("relatorio", "")
        erros = contexto["form"].errors.as_text() if contexto["form"].errors else ""
        return HttpResponse(f"{relatorio}\n{erros}")

    @staticmethod
    def _redirecionar(*_args, **_kwargs):
        return HttpResponseRedirect("/admin/")

    @patch("apps.pca.importacao.views.render", side_effect=_renderizar.__func__)
    def test_rejeita_extensao_antes_de_processar(self, _renderizar):
        self.client.force_login(self.superuser)

        with patch("apps.pca.importacao.views.executar_import") as executar_import:
            resposta = self.client.post(
                "/admin/importar/",
                {"arquivo": self._arquivo("planilha.txt"), "exercicio": self.exercicio.pk},
            )

        self.assertContains(resposta, "Envie um arquivo .xlsx.")
        executar_import.assert_not_called()

    @patch("apps.pca.importacao.views.render", side_effect=_renderizar.__func__)
    def test_rejeita_arquivo_maior_que_o_limite_antes_de_processar(self, _renderizar):
        self.client.force_login(self.superuser)
        arquivo_grande = self._arquivo(conteudo=b"x" * (10 * 1024 * 1024 + 1))

        with patch("apps.pca.importacao.views.executar_import") as executar_import:
            resposta = self.client.post(
                "/admin/importar/",
                {"arquivo": arquivo_grande, "exercicio": self.exercicio.pk},
            )

        self.assertContains(resposta, "10 MB")
        executar_import.assert_not_called()

    @patch("apps.pca.importacao.views.render", side_effect=_renderizar.__func__)
    def test_upload_valido_exibe_relatorio_sem_persistir(self, _renderizar):
        self.client.force_login(self.superuser)

        resposta = self.client.post(
            "/admin/importar/",
            {"arquivo": self._arquivo(), "exercicio": self.exercicio.pk},
        )

        self.assertContains(resposta, "Processos criados: 30")
        self.assertEqual(Processo.objects.count(), 0)
        self.assertEqual(Acompanhamento.objects.count(), 0)

    def test_template_real_renderiza_upload_e_confirmacao(self):
        self.client.force_login(self.superuser)

        resposta = self.client.post(
            "/admin/importar/",
            {"arquivo": self._arquivo(), "exercicio": self.exercicio.pk},
        )

        self.assertTemplateUsed(resposta, "admin/importacao/importar.html")
        self.assertContains(resposta, "Relatório de conferência")
        self.assertContains(resposta, "Confirmar e gravar")
        self.assertContains(resposta, 'enctype="multipart/form-data"', count=2)
        self.assertContains(resposta, 'name="referencia"')
        self.assertNotContains(resposta, "conteudo_base64")
        self.assertNotContains(resposta, "arquivo_nome")

    def _referencia_da_sessao(self):
        return self.client.session[views.CHAVE_SESSAO_UPLOAD_PENDENTE]["referencia"]

    def _enviar_upload(self, arquivo=None):
        return self.client.post(
            "/admin/importar/",
            {"arquivo": arquivo or self._arquivo(), "exercicio": self.exercicio.pk},
        )

    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    def test_confirmacao_persiste_e_separa_as_autorias(self, _redirecionar):
        self.client.force_login(self.superuser)

        upload = self._enviar_upload()
        self.assertEqual(upload.status_code, 200)
        resposta = self.client.post("/admin/importar/confirmar/", {"referencia": self._referencia_da_sessao()})

        self.assertEqual(resposta.status_code, 302)
        self.assertEqual(Processo.objects.count(), 30)
        self.assertEqual(Acompanhamento.objects.count(), 75)
        evento = EventoImportacao.objects.get()
        self.assertEqual(evento.disparado_por, self.superuser)
        self.assertEqual(evento.processos_criados, 30)
        self.assertEqual(evento.acompanhamentos_criados, 75)
        usuario_importador = get_user_model().objects.get(email="importador@pca.local")
        self.assertTrue(
            HistoricalProcesso.objects.filter(history_user=usuario_importador).exists()
        )
        self.assertFalse(
            HistoricalProcesso.objects.filter(history_user=self.superuser).exists()
        )

    @override_settings(DATA_UPLOAD_MAX_MEMORY_SIZE=2_621_440)
    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    @patch("apps.pca.importacao.views.executar_import")
    def test_confirmacao_aceita_planilha_no_limite_de_dez_megabytes(
        self, executar_import, _redirecionar
    ):
        """A confirmação contém só uma referência, mesmo para arquivo de 10 MiB."""
        self.client.force_login(self.superuser)
        conteudo_no_limite = b"x" * (10 * 1024 * 1024)
        relatorio = SimpleNamespace(
            processos_criados=0, acompanhamentos_criados=0,
            caminho_relatorio=Path("/tmp/relatorio-importacao.txt"), formatar=lambda: "ok",
        )
        caminhos = []

        def _executar(caminho, _usuario, _exercicio, *, dry_run):
            caminhos.append((Path(caminho), dry_run))
            self.assertTrue(Path(caminho).is_file())
            self.assertEqual(Path(caminho).stat().st_size, len(conteudo_no_limite))
            return relatorio

        executar_import.side_effect = _executar
        with tempfile.TemporaryDirectory() as diretorio, patch.object(
            views, "DIRETORIO_UPLOAD_PENDENTE", Path(diretorio)
        ):
            upload = self._enviar_upload(self._arquivo("no-limite.xlsx", conteudo_no_limite))
            self.assertEqual(upload.status_code, 200)
            referencia = self._referencia_da_sessao()
            self.assertLess(len(referencia), 100)
            self.assertNotContains(upload, "no-limite.xlsx")
            self.assertNotContains(upload, 'name="arquivo_nome"')
            self.assertNotContains(upload, 'name="conteudo_base64"')
            self.assertNotContains(upload, 'name="exercicio" value=')
            resposta = self.client.post("/admin/importar/confirmar/", {"referencia": referencia})

        self.assertEqual(resposta.status_code, 302)
        self.assertEqual([dry_run for _caminho, dry_run in caminhos], [True, False])
        self.assertTrue(caminhos[0][0].name.endswith(views.SUFIXO_PENDING))
        self.assertTrue(caminhos[1][0].name.endswith(views.SUFIXO_CLAIMED))
        self.assertFalse(caminhos[1][0].exists())
        self.assertNotIn(views.CHAVE_SESSAO_UPLOAD_PENDENTE, self.client.session)

    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    @patch("apps.pca.importacao.views.executar_import")
    def test_referencia_forjada_ou_repetida_nao_executa_importacao(self, executar_import, _redirecionar):
        self.client.force_login(self.superuser)
        relatorio = SimpleNamespace(
            processos_criados=0, acompanhamentos_criados=0,
            caminho_relatorio=Path("/tmp/relatorio-importacao.txt"), formatar=lambda: "ok",
        )
        executar_import.return_value = relatorio
        with tempfile.TemporaryDirectory() as diretorio, patch.object(
            views, "DIRETORIO_UPLOAD_PENDENTE", Path(diretorio)
        ):
            self._enviar_upload()
            referencia = self._referencia_da_sessao()
            forjada = f"{'A' if referencia[0] != 'A' else 'B'}{referencia[1:]}"
            resposta_forjada = self.client.post("/admin/importar/confirmar/", {"referencia": forjada})
            self.assertEqual(resposta_forjada.status_code, 302)
            executar_import.assert_called_once()  # somente o dry-run
            resposta_valida = self.client.post("/admin/importar/confirmar/", {"referencia": referencia})
            self.assertEqual(resposta_valida.status_code, 302)
            resposta_replay = self.client.post("/admin/importar/confirmar/", {"referencia": referencia})
            self.assertEqual(resposta_replay.status_code, 302)
        self.assertEqual(executar_import.call_count, 2)
        self.assertEqual(EventoImportacao.objects.count(), 1)

    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    @patch("apps.pca.importacao.views.executar_import")
    def test_referencia_expirada_descarta_upload_sem_executar_confirmacao(
        self, executar_import, _redirecionar
    ):
        self.client.force_login(self.superuser)
        executar_import.return_value = SimpleNamespace(formatar=lambda: "ok")
        with tempfile.TemporaryDirectory() as diretorio, patch.object(
            views, "DIRETORIO_UPLOAD_PENDENTE", Path(diretorio)
        ):
            self._enviar_upload()
            referencia = self._referencia_da_sessao()
            caminho = views._caminho_upload(referencia, views.SUFIXO_PENDING)
            sessao = self.client.session
            registro = sessao[views.CHAVE_SESSAO_UPLOAD_PENDENTE]
            registro["criado_em"] -= views.TTL_UPLOAD_SEGUNDOS + 1
            sessao[views.CHAVE_SESSAO_UPLOAD_PENDENTE] = registro
            sessao.save()
            resposta = self.client.post("/admin/importar/confirmar/", {"referencia": referencia})
            self.assertEqual(resposta.status_code, 302)
            self.assertFalse(caminho.exists())
            self.assertNotIn(views.CHAVE_SESSAO_UPLOAD_PENDENTE, self.client.session)
        executar_import.assert_called_once_with(
            str(caminho), views._usuario_importador(), self.exercicio, dry_run=True
        )

    @patch("apps.pca.importacao.views.executar_import")
    def test_erro_no_dry_run_descarta_pending_e_sessao(self, executar_import):
        self.client.force_login(self.superuser)
        caminhos = []

        def _falhar(caminho, _usuario, _exercicio, *, dry_run):
            self.assertTrue(dry_run)
            caminhos.append(Path(caminho))
            self.assertTrue(caminhos[0].exists())
            raise CommandError("Planilha inválida")

        executar_import.side_effect = _falhar
        with tempfile.TemporaryDirectory() as diretorio, patch.object(
            views, "DIRETORIO_UPLOAD_PENDENTE", Path(diretorio)
        ):
            resposta = self._enviar_upload()
            self.assertEqual(resposta.status_code, 200)
            self.assertFalse(caminhos[0].exists())
            self.assertEqual(list(Path(diretorio).iterdir()), [])
            self.assertNotIn(views.CHAVE_SESSAO_UPLOAD_PENDENTE, self.client.session)

        self.assertEqual(EventoImportacao.objects.count(), 0)

    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    @patch("apps.pca.importacao.views.executar_import")
    def test_erro_na_confirmacao_descarta_claimed_e_nao_cria_evento(
        self, executar_import, _redirecionar
    ):
        self.client.force_login(self.superuser)
        relatorio = SimpleNamespace(formatar=lambda: "ok")
        caminho_claimed = None

        def _executar(caminho, _usuario, _exercicio, *, dry_run):
            nonlocal caminho_claimed
            if dry_run:
                return relatorio
            caminho_claimed = Path(caminho)
            self.assertTrue(caminho_claimed.exists())
            self.assertTrue(caminho_claimed.name.endswith(views.SUFIXO_CLAIMED))
            raise CommandError("Falha ao gravar")

        executar_import.side_effect = _executar
        with tempfile.TemporaryDirectory() as diretorio, patch.object(
            views, "DIRETORIO_UPLOAD_PENDENTE", Path(diretorio)
        ):
            self._enviar_upload()
            referencia = self._referencia_da_sessao()
            resposta = self.client.post("/admin/importar/confirmar/", {"referencia": referencia})
            self.assertEqual(resposta.status_code, 302)
            self.assertFalse(caminho_claimed.exists())
            self.assertNotIn(views.CHAVE_SESSAO_UPLOAD_PENDENTE, self.client.session)

        self.assertEqual(EventoImportacao.objects.count(), 0)

    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    @patch("apps.pca.importacao.views.executar_import")
    def test_segundo_upload_descarta_pending_anterior_e_so_novo_token_confirma(
        self, executar_import, _redirecionar
    ):
        self.client.force_login(self.superuser)
        relatorio = SimpleNamespace(
            processos_criados=0,
            acompanhamentos_criados=0,
            caminho_relatorio=Path("/tmp/relatorio-importacao.txt"),
            formatar=lambda: "ok",
        )
        executar_import.return_value = relatorio
        with tempfile.TemporaryDirectory() as diretorio, patch.object(
            views, "DIRETORIO_UPLOAD_PENDENTE", Path(diretorio)
        ):
            self._enviar_upload(self._arquivo("primeira.xlsx"))
            referencia_anterior = self._referencia_da_sessao()
            pendente_anterior = views._caminho_upload(
                referencia_anterior, views.SUFIXO_PENDING
            )
            self.assertTrue(pendente_anterior.exists())

            self._enviar_upload(self._arquivo("segunda.xlsx"))
            referencia_atual = self._referencia_da_sessao()
            self.assertNotEqual(referencia_anterior, referencia_atual)
            self.assertFalse(pendente_anterior.exists())
            self.assertTrue(
                views._caminho_upload(referencia_atual, views.SUFIXO_PENDING).exists()
            )

            resposta_anterior = self.client.post(
                "/admin/importar/confirmar/", {"referencia": referencia_anterior}
            )
            self.assertEqual(resposta_anterior.status_code, 302)
            self.assertEqual(executar_import.call_count, 2)

            resposta_atual = self.client.post(
                "/admin/importar/confirmar/", {"referencia": referencia_atual}
            )
            self.assertEqual(resposta_atual.status_code, 302)

        self.assertEqual(executar_import.call_count, 3)
        self.assertEqual(EventoImportacao.objects.count(), 1)

    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    @patch("apps.pca.importacao.views.executar_import")
    def test_ttl_remove_pending_mas_nunca_claimed_em_execucao(self, executar_import, _redirecionar):
        self.client.force_login(self.superuser)
        relatorio = SimpleNamespace(
            processos_criados=0, acompanhamentos_criados=0,
            caminho_relatorio=Path("/tmp/relatorio-importacao.txt"), formatar=lambda: "ok",
        )
        with tempfile.TemporaryDirectory() as diretorio, patch.object(
            views, "DIRETORIO_UPLOAD_PENDENTE", Path(diretorio)
        ):
            def _executar(caminho, _usuario, _exercicio, *, dry_run):
                caminho = Path(caminho)
                if not dry_run:
                    self.assertTrue(caminho.name.endswith(views.SUFIXO_CLAIMED))
                    os.utime(caminho, (0, 0))
                    views.limpar_uploads_pendentes_expirados(agora=views.TTL_UPLOAD_SEGUNDOS + 1)
                    self.assertTrue(caminho.exists())
                    self.assertEqual(caminho.read_bytes(), self.conteudo)
                return relatorio

            executar_import.side_effect = _executar
            self._enviar_upload()
            referencia = self._referencia_da_sessao()
            pendente = views._caminho_upload(referencia, views.SUFIXO_PENDING)
            os.utime(pendente, (0, 0))
            outro_pendente = views._caminho_upload("A" * views.COMPRIMENTO_REFERENCIA, views.SUFIXO_PENDING)
            outro_pendente.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
            outro_pendente.write_bytes(b"expirado")
            os.utime(outro_pendente, (0, 0))
            resposta = self.client.post("/admin/importar/confirmar/", {"referencia": referencia})
            self.assertEqual(resposta.status_code, 302)
            self.assertFalse(outro_pendente.exists())
            self.assertFalse(views._caminho_upload(referencia, views.SUFIXO_CLAIMED).exists())

    def test_staff_sem_superuser_nunca_alcanca_o_importador(self):
        self.client.force_login(self.staff)

        with patch("apps.pca.importacao.views.executar_import") as executar_import:
            for caminho in ("/admin/importar/", "/admin/importar/confirmar/"):
                for metodo in (self.client.get, self.client.post):
                    resposta = metodo(caminho)
                    self.assertIn(resposta.status_code, (302, 403))

        executar_import.assert_not_called()

    @patch("apps.pca.importacao.views.redirect", side_effect=_redirecionar.__func__)
    def test_confirmacao_direta_nao_persiste_nada(self, _redirecionar):
        self.client.force_login(self.superuser)

        resposta = self.client.post("/admin/importar/confirmar/", {})

        self.assertEqual(resposta.status_code, 302)
        self.assertEqual(Processo.objects.count(), 0)
        self.assertEqual(Acompanhamento.objects.count(), 0)
        self.assertEqual(EventoImportacao.objects.count(), 0)

    def test_download_do_modelo_entrega_contrato_vazio_e_catalogos(self):
        """O arquivo novo é estruturalmente idêntico ao contrato do importador."""
        Categoria.objects.create(nome="=catálogo literal")
        self.client.force_login(self.superuser)

        resposta = self.client.get(reverse("admin:pca_eventoimportacao_modelo"))

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(
            resposta["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="modelo-importacao-pca.xlsx"', resposta["Content-Disposition"])
        self.assertEqual(resposta["X-Content-Type-Options"], "nosniff")
        self.assertEqual(int(resposta["Content-Length"]), len(resposta.content))

        workbook = load_workbook(BytesIO(resposta.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Planilha1", "Acompanhamento", "Listas"])
        planilha, acompanhamento, listas = workbook.worksheets
        self.assertEqual(planilha.max_column, 40)
        self.assertEqual(acompanhamento.max_column, 11)
        self.assertEqual(planilha.max_row, 1)
        self.assertEqual(acompanhamento.max_row, 1)
        self.assertEqual(planilha.freeze_panes, "A2")
        self.assertEqual(acompanhamento.freeze_panes, "A2")
        self.assertEqual(listas.freeze_panes, "A2")
        self.assertEqual(planilha.auto_filter.ref, "A1:AN1")
        self.assertEqual(acompanhamento.auto_filter.ref, "A1:K1")
        self.assertTrue(listas.auto_filter.ref.startswith("A1:"))
        self.assertEqual(
            [cell.value for cell in planilha[1]][:3],
            ["Item PCA (ID)", "Nº do Processo SEI", "Descrição do Objeto"],
        )
        self.assertEqual(
            [cell.value for cell in acompanhamento[1]],
            [
                "ID Registro", "ID Processo", "Nº do Processo SEI", "Data da Reunião",
                "Evento", "Tipo de Evento", "Situação Informada (original)",
                "Situação Normalizada", "Prazo Prometido", "Data do Evento Informada",
                "Área Informada",
            ],
        )
        cabecalho_listas = [cell.value for cell in listas[1]]
        self.assertEqual(cabecalho_listas[: len(COLUNAS_CATALOGO)], [nome for nome, _ in COLUNAS_CATALOGO])
        for indice, (_nome, modelo) in enumerate(COLUNAS_CATALOGO, start=1):
            valores = [
                listas.cell(row=row, column=indice).value
                for row in range(2, listas.max_row + 1)
                if listas.cell(row=row, column=indice).value
            ]
            self.assertEqual(
                valores,
                list(modelo.objects.order_by(*modelo._meta.ordering).values_list("nome", flat=True)),
            )
        celula_formula = next(cell for cell in listas["B"] if cell.value == "=catálogo literal")
        self.assertEqual(celula_formula.data_type, "s")
        self.assertFalse(celula_formula.value.startswith("'"))
        self.assertGreater(len(planilha.data_validations.dataValidation), 0)
        self.assertGreater(len(acompanhamento.data_validations.dataValidation), 0)

    def test_modelo_baixado_tem_prazo_e_nao_tem_justificativa_email(self):
        # O modelo distribuído para download segue o mesmo contrato do
        # importador: "Data prevista para entrega" na posição 14, "Data de
        # Recebimento do Processo no Gelic" logo depois, e a justificativa
        # por e-mail fora das duas abas.
        self.client.force_login(self.superuser)

        resposta = self.client.get(reverse("admin:pca_eventoimportacao_modelo"))

        workbook = load_workbook(BytesIO(resposta.content), data_only=False)
        planilha = workbook["Planilha1"]
        cabecalho = [cell.value for cell in planilha[1]]
        self.assertEqual(cabecalho[13], "Data prevista para entrega")
        self.assertEqual(
            cabecalho[14], "Data de Recebimento do Processo no Gelic"
        )
        self.assertNotIn("Justificativa de Alteração Enviada por E-mail", cabecalho)
        self.assertEqual(len(cabecalho), 40)

        acompanhamento = workbook["Acompanhamento"]
        cabecalho_acompanhamento = [cell.value for cell in acompanhamento[1]]
        self.assertNotIn(
            "Justificativa de Alteração Enviada por E-mail", cabecalho_acompanhamento
        )
        self.assertEqual(len(cabecalho_acompanhamento), 11)

        listas = workbook["Listas"]
        cabecalho_listas = [cell.value for cell in listas[1]]
        self.assertNotIn("Justificativa por E-mail", cabecalho_listas)

    def test_modelo_reenviado_faz_dry_run_sem_persistir(self):
        self.client.force_login(self.superuser)
        download = self.client.get(reverse("admin:pca_eventoimportacao_modelo"))

        resposta = self.client.post(
            "/admin/importar/",
            {
                "arquivo": SimpleUploadedFile(
                    "modelo-importacao-pca.xlsx",
                    download.content,
                    content_type=download["Content-Type"],
                ),
                "exercicio": self.exercicio.pk,
            },
        )

        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Processos criados: 0")
        self.assertEqual(Processo.objects.count(), 0)
        self.assertEqual(Acompanhamento.objects.count(), 0)
        self.assertEqual(EventoImportacao.objects.count(), 0)

    def test_download_do_modelo_exige_superuser_e_aceita_apenas_get(self):
        url = reverse("admin:pca_eventoimportacao_modelo")
        self.client.force_login(self.staff)
        resposta_staff = self.client.get(url)
        self.assertIn(resposta_staff.status_code, (302, 403))

        self.client.force_login(self.superuser)
        resposta_post = self.client.post(url)
        self.assertEqual(resposta_post.status_code, 405)
        self.assertEqual(EventoImportacao.objects.count(), 0)

    def test_tela_de_importacao_exibe_link_para_baixar_modelo(self):
        self.client.force_login(self.superuser)

        resposta = self.client.get("/admin/importar/")

        self.assertContains(resposta, "Baixar modelo de planilha")
        self.assertContains(resposta, reverse("admin:pca_eventoimportacao_modelo"))
