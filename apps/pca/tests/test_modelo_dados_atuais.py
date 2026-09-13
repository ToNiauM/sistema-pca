"""Teste de ida-e-volta: prova que o modelo pré-preenchido gerado
(`ops/relatorios/dados-atuais/modelo-pca-dados-atuais-*.xlsx`, via
`manage.py gerar_modelo_dados_atuais`) é aceito pelo importador sem
nenhuma edição, contra o banco de teste (nunca produção) — provando que o
contrato de colunas de `gerar_modelo_preenchido` casa com o que
`ImportadorPca._importar_processos`/`_importar_acompanhamentos` leem hoje.
"""

import glob
import os
from collections import Counter

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TransactionTestCase

from apps.catalogo.models import Exercicio
from apps.pca.importacao.executor import USUARIO_PADRAO, executar_import
from apps.pca.importacao.parsers import parse_data_excel
from apps.pca.models import Acompanhamento, Processo
from apps.pca.utils import calcular_origem_hash

DIRETORIO_DADOS_ATUAIS = "ops/relatorios/dados-atuais"
PADRAO_ARQUIVO = "modelo-pca-dados-atuais-*.xlsx"


def _arquivo_mais_recente():
    """Localiza o `.xlsx` mais recente gerado pela Task 1. Retorna `None`
    quando o artefato manual não existe — o chamador decide entre pular o
    teste (clone limpo, CI) ou falhar (quem já gerou o arquivo e espera o
    round-trip)."""
    candidatos = glob.glob(os.path.join(DIRETORIO_DADOS_ATUAIS, PADRAO_ARQUIVO))
    if not candidatos:
        return None
    return max(candidatos, key=os.path.getmtime)


class TestRoundTripModeloDadosAtuais(TransactionTestCase):
    """Round-trip export -> import inteiramente contra o banco de teste
    (criado/destruído por `manage.py test`), nunca produção.
    `TransactionTestCase` (não `TestCase`): `executar_import` gerencia sua
    própria transação (`@transaction.atomic`), mesmo motivo de
    `TestImportComando` (test_import.py). `serialized_rollback = True`
    preserva entre testes o `Exercicio(ano=2026)` e o usuário
    `importador@pca.local` que nascem por migração de dados — sem eles,
    `executar_import` recusaria a importação."""

    serialized_rollback = True

    def test_round_trip_modelo_dados_atuais_importa_sem_erro_com_contagens_identicas(self):
        caminho = _arquivo_mais_recente()
        if caminho is None:
            self.skipTest(
                f"Nenhum arquivo '{PADRAO_ARQUIVO}' em '{DIRETORIO_DADOS_ATUAIS}/' — "
                "artefato manual, gitignored, não existe num clone limpo. Para gerá-lo: "
                "`manage.py gerar_modelo_dados_atuais > "
                f"{DIRETORIO_DADOS_ATUAIS}/modelo-pca-dados-atuais-$(date +%Y%m%d-%H%M).xlsx`."
            )

        # Tally esperada: lida do PRÓPRIO ARQUIVO via openpyxl, nunca do
        # banco — é a planilha que faz o papel de "verdade externa" que o
        # usuário vai conferir.
        workbook = openpyxl.load_workbook(caminho, data_only=True)
        linhas_planilha1 = [
            linha
            for linha in workbook["Planilha1"].iter_rows(min_row=2, values_only=True)
            if linha[0] is not None  # pula rodapé em branco (mesmo critério do executor)
        ]
        linhas_acompanhamento = [
            linha
            for linha in workbook["Acompanhamento"].iter_rows(min_row=2, values_only=True)
            if linha[1] is not None  # ID Processo — mesmo critério do executor
        ]
        workbook.close()

        tally_uo_arquivo = Counter(linha[6] for linha in linhas_planilha1)  # UO
        tally_tipo_arquivo = Counter(linha[4] for linha in linhas_planilha1)  # Tipo
        total_processos_arquivo = len(linhas_planilha1)
        total_linhas_acompanhamento_arquivo = len(linhas_acompanhamento)

        self.assertGreater(
            total_processos_arquivo, 0,
            "Arquivo gerado sem nenhuma linha de dados em Planilha1 — nada a provar.",
        )

        exercicio = Exercicio.objects.get(ano=2026)
        usuario = get_user_model().objects.get(email=USUARIO_PADRAO)

        # A chave natural do importador (ano+item_pca+referencia_data+
        # situacao_informada, `calcular_origem_hash`) não é única nos
        # dados reais: registros criados pela UI usam hash salgado com
        # UUID e podem coexistir com outro Acompanhamento que, por
        # coincidência de texto, tem a mesma chave natural.
        # `executar_import` deduplica também dentro do próprio lote,
        # então o total esperado pós-import é o nº de hashes distintos no
        # arquivo, não o total de linhas — a diferença é contabilizada
        # como "ignorados" no relatório, nunca perdida em silêncio.
        hashes_arquivo = set()
        for linha in linhas_acompanhamento:
            item_pca = linha[1]
            referencia_data = parse_data_excel(linha[3])
            # `calcular_origem_hash` já normaliza `situacao_informada` via
            # `(valor or "").strip()` internamente — mesmo texto bruto que
            # `_importar_acompanhamentos` usa (`linha[6]`, sem `_texto()`).
            hashes_arquivo.add(
                calcular_origem_hash(exercicio.ano, item_pca, referencia_data, linha[6])
            )
        total_acompanhamentos_arquivo = len(hashes_arquivo)
        duplicatas_intra_arquivo = total_linhas_acompanhamento_arquivo - total_acompanhamentos_arquivo

        # Banco de teste isolado: nada preexistente para este exercício
        # antes do import (garante que a contagem pós-import é 100% do
        # arquivo, não sobra de outro teste).
        self.assertEqual(Processo.objects.filter(exercicio=exercicio).count(), 0)

        relatorio = executar_import(
            caminho, usuario, exercicio, dry_run=False, escrever_stdout=lambda _t: None,
        )

        # Nenhum PROCESSO preexistia no banco de teste vazio — zero
        # ignorados, nenhum CommandError (se tivesse estourado, o teste já
        # teria falhado por exceção antes de chegar aqui). `item_pca` é
        # chave natural única por construção (constraint de banco), sem o
        # problema de colisão do parágrafo acima.
        self.assertEqual(relatorio.processos_ignorados, 0)
        # A diferença entre "linhas no arquivo" e "hashes distintos" tem
        # que aparecer contabilizada no relatório nominal — prova de que
        # nada foi perdido em silêncio, só deduplicado por identidade
        # legítima.
        self.assertEqual(
            relatorio.acompanhamentos_ignorados, duplicatas_intra_arquivo,
            "Nº de acompanhamentos ignorados no relatório diverge do nº de "
            "colisões de chave natural detectadas no arquivo.",
        )

        total_processos_banco = Processo.objects.filter(exercicio=exercicio).count()
        total_acompanhamentos_banco = Acompanhamento.objects.filter(
            processo__exercicio=exercicio
        ).count()

        self.assertEqual(
            total_processos_banco, total_processos_arquivo,
            "Contagem de processos importados diverge das linhas de Planilha1 no arquivo.",
        )
        self.assertEqual(
            total_acompanhamentos_banco, total_acompanhamentos_arquivo,
            "Contagem de acompanhamentos importados diverge dos HASHES DISTINTOS "
            "de Acompanhamento no arquivo (ver nota sobre colisão de chave natural).",
        )

        tally_uo_banco = Counter(
            Processo.objects.filter(exercicio=exercicio).values_list(
                "unidade_organizacional__nome", flat=True
            )
        )
        tally_tipo_banco = Counter(
            Processo.objects.filter(exercicio=exercicio).values_list(
                "tipo__nome", flat=True
            )
        )

        self.assertEqual(
            tally_uo_banco, tally_uo_arquivo,
            "Diff de contagem por UO entre arquivo e banco pós-import não está vazio.",
        )
        self.assertEqual(
            tally_tipo_banco, tally_tipo_arquivo,
            "Diff de contagem por Tipo entre arquivo e banco pós-import não está vazio.",
        )
