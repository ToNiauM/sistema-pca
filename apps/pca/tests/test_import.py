import re
import tempfile
from datetime import date
from io import StringIO
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management import call_command, CommandError
from django.db.models import F
from django.test import SimpleTestCase, TransactionTestCase
from openpyxl import Workbook

from apps.catalogo.models import Categoria, Exercicio, SituacaoExercicio, Tipo, Unidade
from apps.pca.importacao.executor import COLUNAS_CATALOGO, USUARIO_PADRAO
from apps.pca.importacao.modelo import CABECALHOS_ACOMPANHAMENTO, CABECALHOS_AUXILIARES, CABECALHOS_PLANILHA1
from apps.pca.importacao.parsers import (
    parse_booleano_sim,
    parse_mes_previsto,
    parse_numeros_sei,
    parse_status,
)
from apps.pca.models import (
    Acompanhamento,
    Estado,
    HistoricalAcompanhamento,
    HistoricalProcesso,
    Processo,
    Situacao,
    SituacaoSei,
)
from apps.pca.utils import calcular_origem_hash

ARQUIVO_REAL = "apps/pca/fixtures/modelo-controle-exemplo.xlsx"


class TestParsers(SimpleTestCase):
    """Task 1 — funções puras de `apps/pca/importacao/parsers.py`, testadas
    isoladamente (sem ORM), com os valores reais citados no
    01-RESEARCH.md."""

    def test_parse_mes_previsto(self):
        self.assertEqual(parse_mes_previsto("jan"), 1)
        self.assertEqual(parse_mes_previsto("dez"), 12)
        # Nunca inventar mês nem faixa.
        self.assertIsNone(parse_mes_previsto("jan a dez"))

    def test_parse_status_descarta_sufixo_de_prazo(self):
        self.assertEqual(
            parse_status("Não iniciado (dentro do prazo)"), "nao_iniciado"
        )
        self.assertEqual(
            parse_status("Não iniciado (fora do prazo)"), "nao_iniciado"
        )

    def test_parse_booleano_sim(self):
        self.assertTrue(parse_booleano_sim("SIM"))
        self.assertFalse(parse_booleano_sim(""))
        self.assertFalse(parse_booleano_sim(None))

    def test_parse_numeros_sei(self):
        # Célula sintética com três números separados por espaço simples
        # (mesmo formato do item 7 da fixture de exemplo).
        self.assertEqual(
            parse_numeros_sei(
                "12345678000199.000082/2026-53 12345678000199.000083/2026-06 "
                "12345678000199.000089/2026-75"
            ),
            [
                "12345678000199.000082/2026-53",
                "12345678000199.000083/2026-06",
                "12345678000199.000089/2026-75",
            ],
        )
        # Célula sintética com dois números.
        self.assertEqual(
            parse_numeros_sei(
                "12345678000199.000001/2026-75 12345678000199.000010/2026-66"
            ),
            [
                "12345678000199.000001/2026-75",
                "12345678000199.000010/2026-66",
            ],
        )
        self.assertEqual(parse_numeros_sei(""), [])


class TestImportComando(TransactionTestCase):
    """A prova viva contra a fixture anonimizada
    `apps/pca/fixtures/modelo-controle-exemplo.xlsx`.
    `TransactionTestCase` (não `TestCase`) porque `importar_pca` gerencia
    sua própria transação (`@transaction.atomic` + `set_rollback` no
    dry-run) — rodar dentro da transação envolvente do `TestCase`
    mascararia o comportamento real do `--dry-run`.

    `serialized_rollback = True`: sem isto, o `flush` que o Django roda
    entre testes de `TransactionTestCase` apaga também o usuário de
    serviço criado pela migração de dados `core.0003_usuario_importacao`
    — ele só existiria no primeiro teste da classe."""

    serialized_rollback = True

    def test_golden_numbers(self):
        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())

        self.assertEqual(Processo.objects.count(), 30)
        self.assertEqual(Acompanhamento.objects.count(), 75)
        self.assertEqual(Unidade.objects.count(), 5)

        # Golden numbers recalculados empiricamente contra a fixture
        # anonimizada (`gerar_fixture_exemplo`, 30 itens fictícios).
        # `situacao` nunca é hardcodada: é sempre o que
        # `situacao_inicial()` deriva dos 3 fatos gravados em cada item
        # (data_assinatura_contrato/data_recebimento_gelic/prazo_entrega).
        self.assertEqual(Processo.objects.filter(estado=Estado.ATIVO).count(), 25)
        self.assertEqual(Processo.objects.filter(estado=Estado.CANCELADO).count(), 5)
        self.assertEqual(Processo.objects.filter(situacao=Situacao.NO_PRAZO).count(), 7)
        self.assertEqual(Processo.objects.filter(situacao=Situacao.ATRASADO).count(), 7)
        self.assertEqual(
            Processo.objects.filter(situacao=Situacao.EM_TRAMITACAO).count(), 8
        )
        self.assertEqual(
            Processo.objects.filter(situacao=Situacao.CONCLUIDO).count(), 8
        )
        # Todo processo Tipo=Vigente conclui por natureza, sem exceção.
        self.assertEqual(
            Processo.objects.filter(tipo__nome="Vigente")
            .exclude(situacao=Situacao.CONCLUIDO)
            .count(),
            0,
        )

        self.assertEqual(Tipo.objects.count(), 3)
        self.assertEqual(Processo.objects.filter(tipo__nome="Nova Contratação").count(), 14)
        self.assertEqual(Processo.objects.filter(tipo__nome="Renovação").count(), 11)
        self.assertEqual(Processo.objects.filter(tipo__nome="Vigente").count(), 5)

        self.assertEqual(Categoria.objects.count(), 3)
        self.assertEqual(Processo.objects.filter(categoria__nome="Serviços").count(), 19)
        self.assertEqual(
            Processo.objects.filter(categoria__nome="Solução de TIC").count(), 7
        )
        self.assertEqual(Processo.objects.filter(categoria__nome="Material").count(), 4)

        self.assertEqual(Processo.objects.filter(situacao_sei=SituacaoSei.AUTUADO).count(), 18)
        self.assertEqual(
            Processo.objects.filter(situacao_sei=SituacaoSei.A_AUTUAR).count(), 8
        )
        self.assertEqual(
            Processo.objects.filter(situacao_sei=SituacaoSei.NAO_SE_APLICA).count(), 4
        )

    def test_idempotencia(self):
        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())
        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())

        self.assertEqual(Processo.objects.count(), 30)
        self.assertEqual(Acompanhamento.objects.count(), 75)
        self.assertEqual(Unidade.objects.count(), 5)
        self.assertEqual(HistoricalProcesso.objects.count(), 30)
        self.assertEqual(HistoricalAcompanhamento.objects.count(), 75)

    def test_dry_run_nao_persiste(self):
        call_command("importar_pca", ARQUIVO_REAL, "--dry-run", exercicio=2026, stdout=StringIO())

        self.assertEqual(Processo.objects.count(), 0)
        self.assertEqual(Acompanhamento.objects.count(), 0)

    def test_relatorio_nominal(self):
        """A fixture anonimizada nasce de um round-trip pelo ORM
        (`gerar_modelo_preenchido`), então `mes_previsto` sempre grava um
        inteiro válido 1-12 (nunca um texto ambíguo) — a seção 4 do
        relatório (mês ambíguo) fica sempre vazia para esta fixture. As
        seções 1-3 continuam exercitadas de propósito pelo comando
        gerador (`gerar_fixture_exemplo.ITENS`): item 13 tem envio ao
        Gelic depois do recebimento (seção 1), itens 13 e 27 não têm
        grau de prioridade nem classificação (seção 2), itens 7 e 20 têm
        2 números SEI cada (seção 3)."""
        saida = StringIO()
        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=saida)
        texto = saida.getvalue()

        encontrados_envio = sorted(
            int(n) for n in re.findall(r"- item (\d+): envio", texto)
        )
        self.assertEqual(encontrados_envio, [13])

        encontrados_sem_prioridade = sorted(
            int(n) for n in re.findall(r"- item (\d+)$", texto, re.MULTILINE)
        )
        self.assertEqual(encontrados_sem_prioridade, [13, 27])

        encontrados_multiplos_sei = sorted(
            int(n) for n in re.findall(r"- item (\d+): 12345678000199", texto)
        )
        self.assertEqual(encontrados_multiplos_sei, [7, 20])

        self.assertIn(
            "## 4. Mês previsto ambíguo (gravado como nulo, nunca "
            "inventado) (0)",
            texto,
        )

    def test_history_user_nunca_nulo(self):
        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())

        self.assertEqual(
            HistoricalProcesso.objects.filter(history_user__isnull=True).count(), 0
        )
        usuario = get_user_model().objects.get(email="importador@pca.local")
        primeiro = HistoricalProcesso.objects.first()
        self.assertEqual(primeiro.history_user, usuario)

    def test_status_com_sufixo_nao_quebra_validacao_legado(self):
        # `status` não é mais gravado, mas o sufixo de prazo ("(dentro
        # do prazo)"/"(fora do prazo)") continua tendo que ser removido
        # antes de validar contra os 5 rótulos legados reconhecidos
        # (`STATUS_LEGADO_VALIDOS`) — se `parse_status` parasse de
        # descartar o sufixo, toda linha "Não iniciado (...)" da planilha
        # real estouraria `CommandError` antes de criar o primeiro
        # processo.
        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())

        self.assertEqual(Processo.objects.count(), 30)

    def test_recuperacao_exige_exercicio_existente(self):
        with self.assertRaisesMessage(CommandError, "--exercicio"):
            call_command("importar_pca", ARQUIVO_REAL, stdout=StringIO())

        with self.assertRaisesMessage(CommandError, "Exercício 2099 não existe"):
            call_command(
                "importar_pca", ARQUIVO_REAL, exercicio=2099, stdout=StringIO()
            )

        self.assertFalse(Exercicio.objects.filter(ano=2099).exists())

    def test_recuperacao_anual_e_idempotente_sem_colisao(self):
        exercicio_2026 = Exercicio.objects.get(ano=2026)
        exercicio_2027 = Exercicio.objects.create(
            ano=2027, rotulo="PCA 2027", situacao=SituacaoExercicio.ABERTO
        )
        processo_2026 = Processo.objects.filter(exercicio=exercicio_2026, item_pca=1).first()
        self.assertIsNone(processo_2026)

        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())
        processo_2027 = Processo.objects.get(exercicio=exercicio_2026, item_pca=1)
        processo_2027.pk = None
        processo_2027.exercicio = exercicio_2027
        processo_2027.descricao_objeto = "Item 1 de 2027"
        processo_2027.save()

        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())

        self.assertEqual(Processo.objects.filter(exercicio=exercicio_2026).count(), 30)
        self.assertEqual(
            Acompanhamento.objects.filter(processo__exercicio=exercicio_2026).count(), 75
        )
        self.assertEqual(Processo.objects.filter(exercicio=exercicio_2027).count(), 1)
        self.assertEqual(
            Processo.objects.get(exercicio=exercicio_2027, item_pca=1).descricao_objeto,
            "Item 1 de 2027",
        )
        self.assertEqual(
            HistoricalProcesso.objects.filter(
                id=processo_2027.pk, history_change_reason="import inicial"
            ).count(),
            0,
        )

        usuario = get_user_model().objects.get(email="importador@pca.local")
        historicos = HistoricalProcesso.objects.filter(
            exercicio_id=exercicio_2026.pk,
            history_user=usuario,
            history_change_reason="import inicial",
        )
        self.assertEqual(historicos.count(), 30)
        self.assertEqual(
            HistoricalAcompanhamento.objects.filter(
                history_user=usuario,
                history_change_reason="import inicial",
            ).count(),
            75,
        )
        self.assertEqual(
            HistoricalProcesso.objects.filter(
                exercicio_id=exercicio_2026.pk,
                history_user=usuario,
                history_change_reason="import inicial",
            ).count(),
            30,
        )

    def test_hash_de_acompanhamento_isola_exercicio(self):
        data = __import__("datetime").date(2026, 7, 28)
        self.assertNotEqual(
            calcular_origem_hash(2026, 1, data, "Em tramitação"),
            calcular_origem_hash(2027, 1, data, "Em tramitação"),
        )

    def test_dry_run_anual_nao_persiste(self):
        Exercicio.objects.get(ano=2026)
        saida = StringIO()
        call_command(
            "importar_pca", ARQUIVO_REAL, exercicio=2026, dry_run=True, stdout=saida
        )

        self.assertIn("dry-run", saida.getvalue())
        self.assertEqual(Processo.objects.count(), 0)
        self.assertEqual(Acompanhamento.objects.count(), 0)


class TestReindexacaoPrazo(TransactionTestCase):
    """O contrato posicional: "Data prevista para entrega" entra na
    posição 14 de `Planilha1`, e "Justificativa de Alteração Enviada por
    E-mail" sai. `CABECALHOS_PLANILHA1` (fonte da verdade do rótulo
    humano) e os índices fixos que `ImportadorPca._importar_processos` lê
    precisam concordar linha a linha — provado aqui com uma linha
    sintética de três datas distintas, sem depender dos 30 processos da
    fixture de exemplo."""

    serialized_rollback = True

    def _construir_workbook_sintetico(self, linha_processo):
        """Monta um `.xlsx` mínimo, compatível com `ImportadorPca`: cabeçalhos
        de `CABECALHOS_PLANILHA1`/`CABECALHOS_ACOMPANHAMENTO` (mesma fonte
        usada pelo modelo de download) e uma aba `Listas` com só os rótulos
        de catálogo exigidos pela linha sintética — nunca uma cópia
        divergente dos cabeçalhos reais."""
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Planilha1"
        for coluna, cabecalho in enumerate(CABECALHOS_PLANILHA1, start=1):
            planilha.cell(row=1, column=coluna, value=cabecalho)
        for coluna, valor in enumerate(linha_processo, start=1):
            planilha.cell(row=2, column=coluna, value=valor)

        acompanhamento = workbook.create_sheet("Acompanhamento")
        for coluna, cabecalho in enumerate(CABECALHOS_ACOMPANHAMENTO, start=1):
            acompanhamento.cell(row=1, column=coluna, value=cabecalho)

        listas = workbook.create_sheet("Listas")
        cabecalhos_listas = tuple(nome for nome, _ in COLUNAS_CATALOGO) + CABECALHOS_AUXILIARES
        for coluna, cabecalho in enumerate(cabecalhos_listas, start=1):
            listas.cell(row=1, column=coluna, value=cabecalho)
        # Só os rótulos que a linha sintética referencia — suficiente para
        # `_upsert_catalogo` resolver as FKs obrigatórias (tipo/categoria/UO).
        listas.cell(row=2, column=1, value="Nova Contratação")  # Tipo de Contratação
        listas.cell(row=2, column=2, value="Serviços")  # Categoria
        listas.cell(row=2, column=3, value="Unidade Sintética")  # UO

        arquivo = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        workbook.save(arquivo.name)
        workbook.close()
        return Path(arquivo.name)

    def test_modelo_e_executor_compartilham_indices_de_tres_datas(self):
        # A ordem humana do cabeçalho é a mesma que
        # `_importar_processos` consome por índice fixo (0-based): 13 é o
        # prazo planejado novo, 14 a entrega real (fato), 15 a previsão
        # de conclusão da contratação, 16 o status — nunca reordenados de
        # novo sem atualizar os dois lados juntos.
        self.assertEqual(CABECALHOS_PLANILHA1[13], "Data prevista para entrega")
        self.assertEqual(
            CABECALHOS_PLANILHA1[14], "Data de Recebimento do Processo no Gelic"
        )
        self.assertEqual(
            CABECALHOS_PLANILHA1[15], "Data Prevista/Conclusão da Contratação"
        )
        self.assertEqual(CABECALHOS_PLANILHA1[16], "Status")
        self.assertNotIn(
            "Justificativa de Alteração Enviada por E-mail", CABECALHOS_PLANILHA1
        )
        self.assertEqual(len(CABECALHOS_PLANILHA1), 40)

    def test_importa_prazo_sem_reinterpretar_entrega(self):
        # Três datas distintas — se o índice errado for lido, o teste
        # pega a troca de significado, não só a ausência de dado.
        prazo_entrega = date(2026, 5, 15)
        data_recebimento_gelic = date(2026, 6, 20)
        data_prevista_conclusao = date(2026, 7, 25)
        data_envio = date(2026, 4, 1)

        linha = [None] * 40
        linha[0] = 9001  # Item PCA (ID)
        linha[2] = "Objeto sintético"  # Descrição do Objeto
        linha[3] = "Justificativa sintética"
        linha[4] = "Nova Contratação"  # Tipo de Contratação
        linha[5] = "Serviços"  # Categoria
        linha[6] = "Unidade Sintética"  # UO
        linha[7] = 1000  # Valor Estimado
        linha[12] = data_envio  # Data Envio ao Gelic
        linha[13] = prazo_entrega  # Data prevista para entrega (NOVA)
        linha[14] = data_recebimento_gelic  # Data de Recebimento no Gelic (fato)
        linha[15] = data_prevista_conclusao  # Data Prevista/Conclusão
        linha[16] = "Não iniciado"  # Status

        caminho = self._construir_workbook_sintetico(linha)
        try:
            exercicio = Exercicio.objects.get(ano=2026)
            usuario = get_user_model().objects.get(email=USUARIO_PADRAO)
            call_command(
                "importar_pca", str(caminho), exercicio=exercicio.ano,
                usuario=usuario.email, stdout=StringIO(),
            )
        finally:
            caminho.unlink(missing_ok=True)

        processo = Processo.objects.get(exercicio=exercicio, item_pca=9001)
        self.assertEqual(processo.prazo_entrega, prazo_entrega)
        self.assertEqual(processo.data_recebimento_gelic, data_recebimento_gelic)
        self.assertEqual(processo.data_prevista_conclusao, data_prevista_conclusao)
        # As três datas são realmente distintas — nenhuma reaproveitada de
        # outra coluna por engano.
        self.assertNotEqual(processo.prazo_entrega, processo.data_recebimento_gelic)
        self.assertNotEqual(
            processo.data_recebimento_gelic, processo.data_prevista_conclusao
        )

    def test_fixture_de_exemplo_preserva_a_anomalia_de_reindex(self):
        # A mesma checagem de sempre (envio × entrega real, nunca o
        # prazo planejado novo), lida pelos índices reindexados, contra a
        # fixture anonimizada: `gerar_fixture_exemplo` grava de propósito
        # 1 anomalia (item 13, ver seu módulo).
        call_command("importar_pca", ARQUIVO_REAL, exercicio=2026, stdout=StringIO())

        anomalos = Processo.objects.filter(
            data_envio_gelic__isnull=False,
            data_recebimento_gelic__isnull=False,
            data_recebimento_gelic__lt=F("data_envio_gelic"),
        ).count()
        self.assertEqual(anomalos, 1)

        # Ao contrário da planilha real (coluna sempre vazia, zero
        # backfill), a fixture de exemplo popula `prazo_entrega` de
        # propósito no bloco ATRASADO (7 itens,
        # `gerar_fixture_exemplo._PRAZO_ATRASADO`) — é assim que a
        # fixture exercita `situacao_inicial()` caindo no fallback de
        # calendário.
        self.assertEqual(Processo.objects.filter(prazo_entrega__isnull=False).count(), 7)

    def _tentar_importar_status(self, texto_status):
        """Monta uma planilha sintética de uma linha só com `texto_status`
        na coluna Status e roda o import — devolve o `item_pca` usado,
        para os testes de rejeição verificarem a mensagem/o banco."""
        linha = [None] * 40
        linha[0] = 9002  # Item PCA (ID)
        linha[2] = "Objeto sintético"
        linha[3] = "Justificativa sintética"
        linha[4] = "Nova Contratação"
        linha[5] = "Serviços"
        linha[6] = "Unidade Sintética"
        linha[7] = 1000
        linha[16] = texto_status  # Status

        caminho = self._construir_workbook_sintetico(linha)
        try:
            exercicio = Exercicio.objects.get(ano=2026)
            usuario = get_user_model().objects.get(email=USUARIO_PADRAO)
            call_command(
                "importar_pca", str(caminho), exercicio=exercicio.ano,
                usuario=usuario.email, stdout=StringIO(),
            )
        finally:
            caminho.unlink(missing_ok=True)

    def test_status_concluido_nao_contratado_e_rejeitado_explicitamente(self):
        # "finalizado" saiu do domínio; o importador não mapeia o texto
        # cru pra outro status, ele recusa a linha inteira.
        with self.assertRaisesMessage(CommandError, "não é mais suportado"):
            self._tentar_importar_status("Concluído não contratado")
        self.assertEqual(Processo.objects.filter(item_pca=9002).count(), 0)

    def test_status_aguardando_dfd_e_rejeitado_explicitamente(self):
        # Mesma consequência para o outro status removido do domínio.
        with self.assertRaisesMessage(CommandError, "não é mais suportado"):
            self._tentar_importar_status("Aguardando DFD")
        self.assertEqual(Processo.objects.filter(item_pca=9002).count(), 0)

    def test_status_valido_continua_importando_normalmente(self):
        # Confirma que a guarda nova não pega os 5 status válidos junto.
        # "Cancelado" vira `Estado.CANCELADO` (não mais `Processo.status`,
        # que o executor não escreve mais).
        self._tentar_importar_status("Cancelado")
        self.assertEqual(
            Processo.objects.get(item_pca=9002).estado, Estado.CANCELADO
        )


class TestImportEstadoSituacao(TransactionTestCase):
    """`importar_pca` grava `estado`/`situacao` via a mesma função pura
    `situacao_inicial` que o backfill da migração 0023 usa, nunca uma
    segunda lógica de mapeamento divergente."""

    serialized_rollback = True

    def _construir_workbook_sintetico(self, linha_processo, linhas_acompanhamento=None):
        """Cópia do helper de `TestReindexacaoPrazo` (mesma forma, classe
        isolada).

        `linhas_acompanhamento` (opcional) é uma lista de tuplas
        `(item_processo, referencia_data, prazo_prometido,
        situacao_normalizada)` populando a aba Acompanhamento, para os
        testes de promessa vigente na importação. Quando presente,
        também cadastra `situacao_normalizada` na aba Listas (obrigatória
        para `_importar_acompanhamentos`)."""
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Planilha1"
        for coluna, cabecalho in enumerate(CABECALHOS_PLANILHA1, start=1):
            planilha.cell(row=1, column=coluna, value=cabecalho)
        for coluna, valor in enumerate(linha_processo, start=1):
            planilha.cell(row=2, column=coluna, value=valor)

        acompanhamento = workbook.create_sheet("Acompanhamento")
        for coluna, cabecalho in enumerate(CABECALHOS_ACOMPANHAMENTO, start=1):
            acompanhamento.cell(row=1, column=coluna, value=cabecalho)
        for linha_idx, linha_acomp in enumerate(linhas_acompanhamento or (), start=2):
            item_processo, referencia_data, prazo_prometido, situacao_normalizada = (
                linha_acomp
            )
            acompanhamento.cell(row=linha_idx, column=2, value=item_processo)
            acompanhamento.cell(row=linha_idx, column=4, value=referencia_data)
            acompanhamento.cell(row=linha_idx, column=6, value="Reunião de acompanhamento")
            acompanhamento.cell(row=linha_idx, column=7, value=situacao_normalizada)
            acompanhamento.cell(row=linha_idx, column=8, value=situacao_normalizada)
            acompanhamento.cell(row=linha_idx, column=9, value=prazo_prometido)

        listas = workbook.create_sheet("Listas")
        cabecalhos_listas = tuple(nome for nome, _ in COLUNAS_CATALOGO) + CABECALHOS_AUXILIARES
        for coluna, cabecalho in enumerate(cabecalhos_listas, start=1):
            listas.cell(row=1, column=coluna, value=cabecalho)
        listas.cell(row=2, column=1, value="Nova Contratação")  # Tipo de Contratação
        listas.cell(row=3, column=1, value="Vigente")  # Tipo (Teste 2)
        listas.cell(row=2, column=2, value="Serviços")  # Categoria
        listas.cell(row=2, column=3, value="Unidade Sintética")  # UO
        if linhas_acompanhamento:
            idx_situacao = cabecalhos_listas.index("Situação Normalizada") + 1
            for offset, linha_acomp in enumerate(linhas_acompanhamento):
                listas.cell(
                    row=2 + offset, column=idx_situacao, value=linha_acomp[3]
                )

        arquivo = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        workbook.save(arquivo.name)
        workbook.close()
        return Path(arquivo.name)

    def _importar_linha(
        self, item_pca, status_texto, tipo="Nova Contratação", prazo_entrega=None,
        linhas_acompanhamento=None,
    ):
        """Monta e importa uma planilha sintética de uma linha só (mesmo
        padrão de `TestReindexacaoPrazo._tentar_importar_status`), devolvendo
        o `Processo` criado."""
        linha = [None] * 40
        linha[0] = item_pca
        linha[2] = "Objeto sintético"
        linha[3] = "Justificativa sintética"
        linha[4] = tipo
        linha[5] = "Serviços"
        linha[6] = "Unidade Sintética"
        linha[7] = 1000
        linha[13] = prazo_entrega  # Data prevista para entrega
        linha[16] = status_texto  # Status

        caminho = self._construir_workbook_sintetico(linha, linhas_acompanhamento)
        try:
            exercicio = Exercicio.objects.get(ano=2026)
            usuario = get_user_model().objects.get(email=USUARIO_PADRAO)
            call_command(
                "importar_pca", str(caminho), exercicio=exercicio.ano,
                usuario=usuario.email, stdout=StringIO(),
            )
        finally:
            caminho.unlink(missing_ok=True)
        return Processo.objects.get(exercicio=exercicio, item_pca=item_pca)

    def test_teste1_parse_status_inalterado(self):
        # Teste 1 — o parser de célula não muda, só quem consome o
        # resultado (executor.py).
        self.assertEqual(
            parse_status("Concluído (dentro do prazo)"), "concluido"
        )

    def test_teste2_vigente_produz_ativo_concluido(self):
        # Via `situacao_inicial`: Tipo Vigente conclui por natureza,
        # independente de qualquer data.
        processo = self._importar_linha(9101, "Vigente", tipo="Vigente")
        self.assertEqual(processo.estado, Estado.ATIVO)
        self.assertEqual(processo.situacao, Situacao.CONCLUIDO)

    def test_teste3_nao_iniciado_sem_prazo_produz_no_prazo(self):
        # Sem `prazo_entrega`, cai no padrão NO PRAZO.
        processo = self._importar_linha(9102, "Não iniciado", prazo_entrega=None)
        self.assertEqual(processo.situacao, Situacao.NO_PRAZO)

    def test_teste4_nao_iniciado_com_prazo_vencido_produz_atrasado(self):
        # Mesma linha, com `prazo_entrega` no passado.
        processo = self._importar_linha(
            9103, "Não iniciado", prazo_entrega=date(2020, 1, 1)
        )
        self.assertEqual(processo.situacao, Situacao.ATRASADO)

    def test_teste5_cancelado_produz_estado_cancelado(self):
        # Teste 5 — status legado "Cancelado" vira `Estado.CANCELADO`.
        processo = self._importar_linha(9104, "Cancelado")
        self.assertEqual(processo.estado, Estado.CANCELADO)

    def test_teste6_rotulo_desconhecido_continua_estourando_commanderror(self):
        # Regressão: rótulo fora dos 5 reconhecidos continua abortando o
        # import inteiro.
        with self.assertRaisesMessage(CommandError, "não é mais suportado"):
            self._importar_linha(9105, "Rascunho")
        self.assertFalse(Processo.objects.filter(item_pca=9105).exists())

    def test_teste7_promessa_vigente_da_aba_acompanhamento_produz_atrasado(self):
        """Planilha sintética com uma linha em "Acompanhamento" para o
        mesmo item_pca, `prazo_prometido` no passado, e `prazo_entrega`
        (Planilha1) vazio: o processo importado nasce ATRASADO, não
        NO_PRAZO. Cobre o cenário em que a planilha traz o histórico
        completo de Acompanhamento desde o início, mesmo para processos
        "novos" do ponto de vista de um banco vazio."""
        processo = self._importar_linha(
            9106,
            "Não iniciado",
            prazo_entrega=None,
            linhas_acompanhamento=[
                (9106, date(2020, 1, 1), date(2020, 1, 15), "Compromisso de entrega"),
            ],
        )
        self.assertEqual(processo.situacao, Situacao.ATRASADO)
