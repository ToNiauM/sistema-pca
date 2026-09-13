import csv
import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from urllib.parse import parse_qsl, urlparse

from django.contrib.auth import get_user_model
from django.contrib.sessions.backends.db import SessionStore
from django.core.management import call_command
from django.test import RequestFactory, SimpleTestCase, TransactionTestCase
from django.urls import reverse
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.catalogo.models import Categoria, Exercicio, SituacaoExercicio, Tipo, Unidade
from apps.pca.colunas import COLUNAS_PADRAO, canonizar, colunas_selecionadas
from apps.pca.exportacao import COLUNAS, _colunas_resolvidas
from apps.pca.models import Estado, Processo, Situacao

ARQUIVO_REAL = "apps/pca/fixtures/modelo-controle-exemplo.xlsx"

# `COLUNAS` é derivada de `apps.pca.colunas.COLUNAS` (fonte única também
# da tela) + `data_inclusao_pca` (só do arquivo). `CABECALHOS_ESPERADOS` é
# lido da fonte única — nunca retranscrito à mão — para o teste nunca
# divergir por transcrição; os testes abaixo fixam o contrato (rótulos
# com unidade monetária, zero "Status", zero bloco contratual) sobre esse
# valor.
#
# Com todas as 40 chaves selecionadas (`_csv()`/`_xlsx()` abaixo),
# `_colunas_resolvidas` devolve a ordem canônica de `canonizar` — núcleo
# `COLUNAS_PADRAO` primeiro, depois as opcionais na ordem do registro.
# `CABECALHOS_ESPERADOS` reflete essa mesma canonização.
_MAPA_CABECALHO_POR_CHAVE = {chave: cabecalho for chave, cabecalho, _ in COLUNAS}
CABECALHOS_ESPERADOS = [
    _MAPA_CABECALHO_POR_CHAVE[chave]
    for chave in canonizar([chave for chave, _, _ in COLUNAS])
]

# O bloco contratual (colunas 25–39 da planilha) entra por completo no
# registro único como colunas opcionais: os testes afirmam a ausência só
# no padrão compartilhado (`COLUNAS_PADRAO`, sem `?colunas=`), nunca no
# export com todas as chaves selecionadas.

# Posições 1-indexed das colunas tipadas dentro da especificação acima
# (a ordem é a de `apps.pca.colunas.COLUNAS`).
COL_VALOR = CABECALHOS_ESPERADOS.index("Valor previsto (R$)") + 1
COL_VALOR_CONTRATADO = CABECALHOS_ESPERADOS.index("Valor contratado (R$)") + 1
COL_MES = CABECALHOS_ESPERADOS.index("Mês previsto") + 1
# `prazo_vigente`/`prazo_entrega` saíram do registro; `prazo_inicial`
# ("Prazo inicial") entra no lugar de `prazo_entrega` e `prazo_efetivo`
# passa a rotular "Prazo atual".
COL_PRAZO_INICIAL = CABECALHOS_ESPERADOS.index("Prazo inicial") + 1
COL_SEI = CABECALHOS_ESPERADOS.index("Nº do processo SEI") + 1
COL_N_REUNIOES = CABECALHOS_ESPERADOS.index("Nº reuniões") + 1
COL_N_COMPROMISSOS = CABECALHOS_ESPERADOS.index("Nº compromissos") + 1
COL_PRAZO_EFETIVO = CABECALHOS_ESPERADOS.index("Prazo atual") + 1
COL_DIAS_ATRASO = CABECALHOS_ESPERADOS.index("Dias em atraso") + 1
COL_N_PRORROGACOES = CABECALHOS_ESPERADOS.index("Prazos remarcados") + 1
COL_SITUACAO = CABECALHOS_ESPERADOS.index("Situação última reunião") + 1
COL_ESTADO = CABECALHOS_ESPERADOS.index("Estado") + 1
COL_SITUACAO_PROCESSO = CABECALHOS_ESPERADOS.index("Situação") + 1
ULTIMA_COLUNA_LETRA = get_column_letter(len(COLUNAS))

FORMATO_MOEDA = "#,##0.00"
FORMATO_DATA = "dd/mm/yyyy"

# Larguras esperadas (unidades Excel) para as colunas com valor
# fechado; as demais têm só um teto.
LARGURA_ITEM = 10
LARGURA_OBJETO = 60
LARGURA_UNIDADE = 22
LARGURA_VALOR = 22
LARGURA_DATA = 16
LARGURA_SITUACAO = 24


class TestExportacaoArquivos(TransactionTestCase):
    """O contrato pt-BR e tipado dos downloads CSV/XLSX, espelhando
    `apps.pca.colunas` (fonte única).

    Mesmo padrão de `TestFiltrosTabela`/`TestDashboardGoldenNumbers`:
    import real em `setUp()` (`TransactionTestCase` nunca chama
    `setUpTestData()` e faz `flush()` a cada teste), `serialized_rollback`
    para preservar o usuário de serviço da migração de dados. Os testes
    abrem e inspecionam os arquivos (bytes do CSV, células do XLSX via
    openpyxl) — nenhuma inspeção manual no Excel é necessária.
    """

    serialized_rollback = True

    def setUp(self):
        call_command(
            "importar_pca",
            ARQUIVO_REAL,
            "--usuario=importador@pca.local",
            "--exercicio=2026",
            stdout=StringIO(),
        )
        self.usuario = get_user_model().objects.create_user(
            email="leitor@pca.local", password="x-forte-123"
        )
        self.client.force_login(self.usuario)

    # --- helpers ----------------------------------------------------------

    def _csv(self, **params):
        """Baixa o CSV e devolve (resposta, linhas) com as linhas já
        parseadas pelo separador `;` e sem o BOM — o BOM é verificado à
        parte, sobre os bytes crus. Injeta `colunas=` com todas as chaves
        de `exportacao.COLUNAS` como parâmetro repetido — sem isso, uma
        request sem `?colunas=` devolveria só as 5 colunas padrão, e esta
        classe inteira exercita o contrato completo."""
        params.setdefault("colunas", [chave for chave, _, _ in COLUNAS])
        resposta = self.client.get(reverse("pca:exportar_csv"), params)
        linhas = list(
            csv.reader(
                StringIO(resposta.content.decode("utf-8-sig")), delimiter=";"
            )
        )
        return resposta, linhas

    def _xlsx(self, **params):
        """Baixa o XLSX e devolve (resposta, worksheet) com o pacote aberto
        pela openpyxl — prova de que o arquivo é um XLSX válido, não texto
        disfarçado. Mesma injeção de `colunas=` completo do `_csv` acima."""
        params.setdefault("colunas", [chave for chave, _, _ in COLUNAS])
        resposta = self.client.get(reverse("pca:exportar_xlsx"), params)
        planilha = load_workbook(BytesIO(resposta.content))
        return resposta, planilha.active

    def _ids_tabela(self, **params):
        """Percorre TODAS as páginas de /tabela sob o mesmo querystring e
        devolve os item_pca exibidos — mesma técnica de
        `TestFiltrosTabela._todas_paginas` (nunca parsing de HTML)."""
        ids = []
        pagina_num = 1
        while True:
            resposta = self.client.get(
                reverse("pca:tabela"), {**params, "pagina": pagina_num}
            )
            pagina = resposta.context["pagina"]
            ids.extend(p.item_pca for p in pagina.object_list)
            if not pagina.has_next():
                break
            pagina_num += 1
        return ids

    def _linha_csv_do_item(self, linhas, item_pca):
        for linha in linhas[1:]:
            if linha[0] == str(item_pca):
                return linha
        self.fail(f"item {item_pca} não encontrado no CSV exportado")

    def _linha_xlsx_do_item(self, ws, item_pca):
        for linha in range(2, ws.max_row + 1):
            if ws.cell(row=linha, column=1).value == item_pca:
                return linha
        self.fail(f"item {item_pca} não encontrado no XLSX exportado")

    def _filtro_combinado(self):
        """Um querystring combinado e não trivial (0 < N < 30) para provar
        a paridade export × tabela sob o mesmo filtro."""
        alvo = Processo.objects.filter(estado=Estado.CANCELADO.value).first()
        self.assertIsNotNone(alvo)
        params = {
            "estado": Estado.CANCELADO.value,
            "uo": alvo.unidade_organizacional_id,
        }
        esperado = Processo.objects.filter(
            estado=Estado.CANCELADO.value,
            unidade_organizacional_id=alvo.unidade_organizacional_id,
        ).count()
        self.assertGreater(esperado, 0)
        self.assertLess(esperado, 30)
        return params, esperado

    # --- rotas, cabeçalhos HTTP e autenticação -----------------------------

    def test_endpoints_resolvem_por_nome(self):
        self.assertTrue(reverse("pca:exportar_csv").startswith("/exportar"))
        self.assertTrue(reverse("pca:exportar_xlsx").startswith("/exportar"))

    def test_csv_content_type_disposition_e_cache(self):
        resposta, _ = self._csv()
        self.assertEqual(resposta.status_code, 200)
        self.assertTrue(resposta["Content-Type"].startswith("text/csv"))
        self.assertIn("attachment", resposta["Content-Disposition"])
        self.assertIn(".csv", resposta["Content-Disposition"])
        # Download nunca cacheável na borda (dado interno, por sessão
        # autenticada), mesmo contrato das três visões.
        self.assertEqual(resposta["Cache-Control"], "private, no-store")

    def test_xlsx_content_type_disposition_e_cache(self):
        resposta, _ = self._xlsx()
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(
            resposta["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", resposta["Content-Disposition"])
        self.assertIn(".xlsx", resposta["Content-Disposition"])
        self.assertEqual(resposta["Cache-Control"], "private, no-store")

    def test_export_exige_autenticacao(self):
        # O LoginRequiredMiddleware global precisa interceptar os dois
        # endpoints como qualquer outra rota.
        self.client.logout()
        for nome in ("pca:exportar_csv", "pca:exportar_xlsx"):
            with self.subTest(rota=nome):
                resposta = self.client.get(reverse(nome))
                self.assertEqual(resposta.status_code, 302)
                self.assertTrue(resposta.url.startswith("/login/"))

    # --- CSV: estrutura, codificação e tipos -------------------------------

    def test_csv_abre_como_utf8_com_bom(self):
        # Sem o BOM (utf-8-sig), o Excel pt-BR abre o arquivo como
        # Windows-1252 e corrói todos os acentos.
        resposta, _ = self._csv()
        self.assertTrue(resposta.content.startswith(b"\xef\xbb\xbf"))

    def test_csv_cabecalho_ptbr_exato_com_bloco_contratual(self):
        # O bloco de Execução contratual/Publicação faz parte do
        # registro único: com todas as chaves selecionadas (`_csv()`), o
        # cabeçalho inclui essas colunas.
        _, linhas = self._csv()
        self.assertEqual(linhas[0], CABECALHOS_ESPERADOS)
        for esperado in ("Modalidade", "Instrumento contratual", "Data lançamento SPW"):
            self.assertIn(esperado, linhas[0])

    def test_csv_30_linhas_de_dados_sem_filtro(self):
        _, linhas = self._csv()
        # 1 cabeçalho + 30 processos — uma linha por processo; nenhuma
        # linha extra por nº SEI, nenhuma linha de total no CSV.
        self.assertEqual(len(linhas), 31)
        ids = [linha[0] for linha in linhas[1:]]
        self.assertEqual(len(set(ids)), 30)

    def test_csv_decimal_com_virgula_sem_separador_de_milhar(self):
        # Representação CSV (e só ela) usa vírgula decimal, para o Excel
        # pt-BR interpretar a célula como número ao abrir via `;`.
        processo = Processo.objects.exclude(valor_estimado__isnull=True).first()
        self.assertIsNotNone(processo)
        _, linhas = self._csv()
        celula = self._linha_csv_do_item(linhas, processo.item_pca)[COL_VALOR - 1]
        esperado = f"{processo.valor_estimado:.2f}".replace(".", ",")
        self.assertEqual(celula, esperado)
        self.assertRegex(celula, r"^\d+,\d{2}$")

    def test_csv_datas_em_dd_mm_aaaa(self):
        processo = Processo.objects.exclude(data_envio_gelic__isnull=True).first()
        self.assertIsNotNone(processo)
        _, linhas = self._csv()
        linha = self._linha_csv_do_item(linhas, processo.item_pca)
        indice_envio = CABECALHOS_ESPERADOS.index("Envio ao Gelic")
        self.assertEqual(
            linha[indice_envio], processo.data_envio_gelic.strftime("%d/%m/%Y")
        )
        self.assertRegex(linha[indice_envio], r"^\d{2}/\d{2}/\d{4}$")

    def test_csv_ausencia_e_celula_vazia_nao_rotulo(self):
        # O rótulo cinza da tela ("Não classificado", "Sem mês previsto")
        # nunca vai ao arquivo — texto no lugar de número quebra
        # SUM/filtro/ordenação no Excel.
        _, linhas = self._csv()
        sem_mes = Processo.objects.filter(mes_previsto__isnull=True).first()
        self.assertIsNotNone(sem_mes, "fixture de exemplo tem 1 item sem mês previsto")
        linha_sem_mes = self._linha_csv_do_item(linhas, sem_mes.item_pca)
        self.assertEqual(linha_sem_mes[COL_MES - 1], "")

        sem_prioridade = Processo.objects.filter(
            grau_prioridade__isnull=True
        ).first()
        self.assertIsNotNone(sem_prioridade, "fixture de exemplo tem 2 itens sem prioridade")
        linha_np = self._linha_csv_do_item(linhas, sem_prioridade.item_pca)
        self.assertEqual(
            linha_np[CABECALHOS_ESPERADOS.index("Prioridade")], ""
        )
        self.assertEqual(
            linha_np[CABECALHOS_ESPERADOS.index("Classificação")], ""
        )
        # Nenhuma célula do arquivo carrega os rótulos da tela.
        for linha in linhas[1:]:
            for celula in linha:
                self.assertNotIn(celula, ("Não classificado", "Sem mês previsto",
                                          "Não informado", "Sem prazo"))

    def test_csv_multiplos_seis_numa_celula_separados_por_espaco(self):
        # Item 37 tem três nºs SEI, item 72 tem dois; a linha continua
        # uma por processo.
        alvo = None
        for processo in Processo.objects.prefetch_related("numeros_sei"):
            seis = [s.numero_sei for s in processo.numeros_sei.all()]
            if len(seis) >= 2:
                alvo = (processo, seis)
                break
        self.assertIsNotNone(alvo, "dado real tem processos com múltiplos SEIs")
        processo, seis = alvo
        _, linhas = self._csv()
        celula = self._linha_csv_do_item(linhas, processo.item_pca)[COL_SEI - 1]
        self.assertEqual(celula, " ".join(seis))
        self.assertNotIn("  ", celula)  # separador é UM espaço, nunca dois

    def test_csv_derivadas_batem_com_para_listagem(self):
        # As derivadas vêm de graça de para_listagem() e precisam
        # chegar ao arquivo com exatamente os valores da tela — item 17
        # (bloco ATRASADO da fixture de exemplo) tem `prazo_entrega`
        # preenchido, então `prazo_inicial`/`prazo_efetivo` não são nulos.
        esperado = Processo.objects.para_listagem().get(item_pca=17)
        _, linhas = self._csv()
        linha = self._linha_csv_do_item(linhas, 17)
        self.assertEqual(linha[COL_N_REUNIOES - 1], str(esperado.n_reunioes))
        self.assertEqual(linha[COL_N_COMPROMISSOS - 1], str(esperado.n_compromissos))
        # Prazo inicial (COALESCE(prazo_entrega, primeira promessa)) e
        # prazo atual — mesma identidade de sempre.
        self.assertEqual(
            linha[COL_PRAZO_INICIAL - 1], esperado.prazo_inicial.strftime("%d/%m/%Y")
        )
        self.assertEqual(
            linha[COL_PRAZO_EFETIVO - 1], esperado.prazo_efetivo.strftime("%d/%m/%Y")
        )
        self.assertEqual(linha[COL_DIAS_ATRASO - 1], str(esperado.dias_atraso))
        self.assertEqual(linha[COL_N_PRORROGACOES - 1], str(esperado.n_prorrogacoes))
        self.assertEqual(linha[COL_SITUACAO - 1], esperado.situacao_atual)

    def test_csv_processo_com_quatro_promessas_exporta_tres_remarcacoes(self):
        # `n_prorrogacoes` = (nº de `prazo_prometido` distintos no
        # histórico do item) - 1. A fixture de exemplo não tem nenhum
        # item com histórico de remarcação pronto, então o cenário é
        # fabricado aqui: 4 promessas distintas para o item 26 (neutro,
        # sem nenhum fato gravado) -> 3 remarcações.
        from apps.catalogo.models import SituacaoNormalizada
        from apps.pca.models import Acompanhamento, TipoEvento

        processo = Processo.objects.get(item_pca=26)
        situacao_cat = SituacaoNormalizada.objects.first()
        promessas = [date(2026, 2, 1), date(2026, 3, 1), date(2026, 4, 1), date(2026, 5, 1)]
        for indice, prazo in enumerate(promessas):
            Acompanhamento.objects.create(
                processo=processo,
                referencia_data=date(2026, 1, 1) + timedelta(days=indice),
                origem_hash=f"hash-remarcacao-{processo.pk}-{indice}",
                tipo_evento=TipoEvento.REUNIAO_ACOMPANHAMENTO,
                situacao=situacao_cat,
                prazo_prometido=prazo,
            )

        esperado = Processo.objects.para_listagem().get(item_pca=26)
        self.assertEqual(esperado.n_prorrogacoes, 3)
        _, linhas = self._csv()
        linha = self._linha_csv_do_item(linhas, 26)
        self.assertEqual(linha[COL_N_PRORROGACOES - 1], "3")
        self.assertEqual(
            linha[COL_PRAZO_INICIAL - 1], esperado.prazo_inicial.strftime("%d/%m/%Y")
        )

    def test_csv_estado_tipo_situacao_com_valores_corretos(self):
        # "Status" saiu; Estado/Situação/Tipo batem com
        # get_estado_display()/processo.tipo.nome/Situacao(...).label de
        # um processo fixture conhecido. A ordem relativa dessas 3
        # colunas no cabeçalho é a de `apps.pca.colunas` (fonte única da
        # tela): só o conteúdo é fixado aqui, não a posição relativa.
        cabecalho = CABECALHOS_ESPERADOS
        self.assertNotIn("Status", cabecalho)

        processo = Processo.objects.select_related("tipo").get(item_pca=1)
        _, linhas = self._csv()
        linha = self._linha_csv_do_item(linhas, 1)
        self.assertEqual(linha[cabecalho.index("Tipo")], processo.tipo.nome)
        self.assertEqual(
            linha[cabecalho.index("Estado")], processo.get_estado_display()
        )
        # A coluna "Situação" exporta `situacao_efetiva` (a mesma
        # correção-na-leitura que a tela usa via `para_listagem()`), não
        # `get_situacao_display()` cru — exceto quando `estado` é
        # cancelado (ver classe própria abaixo).
        situacao_efetiva = Processo.objects.para_listagem().get(
            item_pca=1
        ).situacao_efetiva
        if processo.estado == Estado.CANCELADO.value:
            self.assertEqual(linha[cabecalho.index("Situação")], "Cancelado")
        else:
            self.assertEqual(
                linha[cabecalho.index("Situação")], Situacao(situacao_efetiva).label
            )

    def test_csv_neutraliza_injecao_de_formula(self):
        # Texto iniciado por =, +, - ou @ vira fórmula ao abrir no Excel;
        # o export prefixa com apóstrofo em todo campo textual,
        # preservando o conteúdo visível.
        base = Processo.objects.first()
        itens = {}
        for indice, prefixo in enumerate(("=", "+", "-", "@"), start=1):
            item = 900 + indice
            Processo.objects.create(
                item_pca=item,
                descricao_objeto=f"{prefixo}HYPERLINK(\"http://x\",\"y\")",
                tipo=base.tipo,
                categoria=base.categoria,
                unidade_organizacional=base.unidade_organizacional,
                exercicio=base.exercicio,
            )
            itens[item] = prefixo
        _, linhas = self._csv()
        indice_objeto = CABECALHOS_ESPERADOS.index("Objeto")
        for item, prefixo in itens.items():
            with self.subTest(prefixo=prefixo):
                celula = self._linha_csv_do_item(linhas, item)[indice_objeto]
                self.assertTrue(celula.startswith("'"))
                self.assertIn(prefixo, celula)

    def test_csv_tem_mesmos_ids_da_tabela_com_querystring_combinado(self):
        # A promessa central: o arquivo é o mesmo conjunto filtrado
        # visível na tela — consumido de queryset_filtrado(request), nunca
        # de um filtro reimplementado, e ignorando a paginação (o CSV
        # leva o conjunto filtrado inteiro).
        params, esperado = self._filtro_combinado()
        ids_tabela = self._ids_tabela(**params)
        self.assertEqual(len(ids_tabela), esperado)
        _, linhas = self._csv(**params)
        ids_csv = [int(linha[0]) for linha in linhas[1:]]
        self.assertEqual(ids_csv, ids_tabela)

    # --- XLSX: tipos nativos, formatos e SUM --------------------------------

    def test_xlsx_cabecalho_identico_ao_csv_e_30_dados(self):
        # Especificação única de colunas: o cabeçalho do XLSX é o mesmo do
        # CSV, sem drift. Região da tabela: 31 linhas incluindo cabeçalho
        # (30 processos); a linha 32 é o total com as fórmulas SUM
        # exigidas pelo plano — fora da região filtrável.
        _, ws = self._xlsx()
        cabecalho = [celula.value for celula in ws[1]]
        self.assertEqual(cabecalho, CABECALHOS_ESPERADOS)
        self.assertEqual(ws.max_row, 32)

    def test_xlsx_valor_e_numero_nativo_com_formato_monetario(self):
        # Pitfall 22 — o valor NUNCA é texto formatado no servidor; é
        # Decimal/número nativo + number_format, para SUM() funcionar.
        processo = Processo.objects.exclude(valor_estimado__isnull=True).first()
        _, ws = self._xlsx()
        linha = self._linha_xlsx_do_item(ws, processo.item_pca)
        celula = ws.cell(row=linha, column=COL_VALOR)
        self.assertIsInstance(celula.value, (int, float, Decimal))
        self.assertNotIsInstance(celula.value, str)
        self.assertAlmostEqual(
            float(celula.value), float(processo.valor_estimado), places=2
        )
        self.assertEqual(celula.number_format, FORMATO_MOEDA)

    def test_xlsx_datas_sao_date_com_formato_dd_mm_aaaa(self):
        processo = Processo.objects.exclude(data_envio_gelic__isnull=True).first()
        _, ws = self._xlsx()
        linha = self._linha_xlsx_do_item(ws, processo.item_pca)
        celula = ws.cell(
            row=linha, column=CABECALHOS_ESPERADOS.index("Envio ao Gelic") + 1
        )
        self.assertIsInstance(celula.value, (datetime, date))
        data = celula.value.date() if isinstance(celula.value, datetime) else celula.value
        self.assertEqual(data, processo.data_envio_gelic)
        self.assertEqual(celula.number_format, FORMATO_DATA)

    def test_xlsx_mes_numerico_e_item_sem_mes_vazio(self):
        # A coluna de mês continua numérica; o único mês nulo da
        # fixture de exemplo é célula vazia de verdade, não o texto "Sem
        # mês previsto" (que rebaixaria a coluna inteira a texto no
        # Excel).
        _, ws = self._xlsx()
        com_mes = Processo.objects.exclude(mes_previsto__isnull=True).first()
        linha = self._linha_xlsx_do_item(ws, com_mes.item_pca)
        celula = ws.cell(row=linha, column=COL_MES)
        self.assertIsInstance(celula.value, int)
        self.assertTrue(1 <= celula.value <= 12)

        sem_mes = Processo.objects.filter(mes_previsto__isnull=True).first()
        self.assertIsNotNone(sem_mes, "fixture de exemplo tem 1 item sem mês previsto")
        linha_sem_mes = self._linha_xlsx_do_item(ws, sem_mes.item_pca)
        self.assertIsNone(ws.cell(row=linha_sem_mes, column=COL_MES).value)

    def test_xlsx_vazios_sao_reais_nao_texto(self):
        # Ausência é célula vazia (None), nunca "" nem o rótulo da tela;
        # o Excel precisa enxergar "(Vazias)" no filtro.
        _, ws = self._xlsx()
        sem_prioridade = Processo.objects.filter(
            grau_prioridade__isnull=True
        ).first()
        linha = self._linha_xlsx_do_item(ws, sem_prioridade.item_pca)
        self.assertIsNone(
            ws.cell(
                row=linha,
                column=CABECALHOS_ESPERADOS.index("Prioridade") + 1,
            ).value
        )
        # E nenhuma célula da coluna de valor é texto (quebraria SUM).
        for r in range(2, ws.max_row):  # exclui a linha de total
            valor = ws.cell(row=r, column=COL_VALOR).value
            self.assertFalse(isinstance(valor, str), f"linha {r} tem valor texto")

    def test_xlsx_formula_sum_calculavel_sobre_valor_previsto_e_contratado(self):
        # As duas colunas monetárias selecionadas (Valor previsto e
        # Valor contratado) ganham `SUM()` real na mesma linha de total,
        # "Total" só na primeira célula.
        _, ws = self._xlsx()
        linha_total = ws.max_row
        self.assertEqual(ws.cell(row=linha_total, column=1).value, "Total")
        letra_valor = get_column_letter(COL_VALOR)
        letra_contratado = get_column_letter(COL_VALOR_CONTRATADO)
        # 30 dados: cabeçalho na linha 1, dados em 2..31, total na 32.
        self.assertEqual(
            ws.cell(row=linha_total, column=COL_VALOR).value,
            f"=SUM({letra_valor}2:{letra_valor}{linha_total - 1})",
        )
        self.assertEqual(
            ws.cell(row=linha_total, column=COL_VALOR_CONTRATADO).value,
            f"=SUM({letra_contratado}2:{letra_contratado}{linha_total - 1})",
        )
        # "Total" só aparece na primeira célula da linha, nunca repetido.
        if COL_VALOR != 1:
            self.assertNotEqual(ws.cell(row=linha_total, column=COL_VALOR).value, "Total")

        # Sob filtro, a fórmula acompanha exatamente a nova região de dados.
        params, esperado = self._filtro_combinado()
        _, ws_filtrado = self._xlsx(**params)
        total_filtrado = ws_filtrado.max_row
        self.assertEqual(total_filtrado, esperado + 2)  # cabeçalho + dados + total
        self.assertEqual(
            ws_filtrado.cell(row=total_filtrado, column=COL_VALOR).value,
            f"=SUM({letra_valor}2:{letra_valor}{total_filtrado - 1})",
        )

    def test_xlsx_autofiltro_e_cabecalho_congelado(self):
        _, ws = self._xlsx()
        self.assertEqual(ws.auto_filter.ref, f"A1:{ULTIMA_COLUNA_LETRA}31")
        self.assertEqual(ws.freeze_panes, "A2")

    def test_xlsx_larguras_de_coluna_batem_com_d_28_35(self):
        # Larguras iniciais: Item 10, Objeto 60, UO 22, valores 22,
        # datas 16, Situação 24.
        _, ws = self._xlsx()
        indice_item = CABECALHOS_ESPERADOS.index("Item") + 1
        indice_objeto = CABECALHOS_ESPERADOS.index("Objeto") + 1
        indice_unidade = CABECALHOS_ESPERADOS.index("Unidade") + 1
        self.assertEqual(
            ws.column_dimensions[get_column_letter(indice_item)].width, LARGURA_ITEM
        )
        self.assertEqual(
            ws.column_dimensions[get_column_letter(indice_objeto)].width,
            LARGURA_OBJETO,
        )
        self.assertEqual(
            ws.column_dimensions[get_column_letter(indice_unidade)].width,
            LARGURA_UNIDADE,
        )
        self.assertEqual(
            ws.column_dimensions[get_column_letter(COL_VALOR)].width, LARGURA_VALOR
        )
        self.assertEqual(
            ws.column_dimensions[get_column_letter(COL_VALOR_CONTRATADO)].width,
            LARGURA_VALOR,
        )
        self.assertEqual(
            ws.column_dimensions[get_column_letter(COL_PRAZO_EFETIVO)].width,
            LARGURA_DATA,
        )
        self.assertEqual(
            ws.column_dimensions[get_column_letter(COL_SITUACAO_PROCESSO)].width,
            LARGURA_SITUACAO,
        )
        # Colunas novas: datas=16, textos=18, moeda=22.
        indice_modalidade = CABECALHOS_ESPERADOS.index("Modalidade") + 1
        indice_data_spw = CABECALHOS_ESPERADOS.index("Data lançamento SPW") + 1
        self.assertEqual(
            ws.column_dimensions[get_column_letter(indice_modalidade)].width, 18
        )
        self.assertEqual(
            ws.column_dimensions[get_column_letter(indice_data_spw)].width, 16
        )

    def test_xlsx_texto_longo_quebra_linha(self):
        # Objeto/Justificativa (TIPO_TEXTO) ganham `wrap_text=True` para
        # textos longos não ficarem cortados.
        _, ws = self._xlsx()
        indice_objeto = CABECALHOS_ESPERADOS.index("Objeto") + 1
        celula = ws.cell(row=2, column=indice_objeto)
        self.assertTrue(celula.alignment.wrap_text)

    def test_xlsx_derivadas_tipadas_batem_com_para_listagem(self):
        esperado = Processo.objects.para_listagem().get(item_pca=1)
        _, ws = self._xlsx()
        linha = self._linha_xlsx_do_item(ws, 1)
        self.assertEqual(
            ws.cell(row=linha, column=COL_N_REUNIOES).value, esperado.n_reunioes
        )
        self.assertEqual(
            ws.cell(row=linha, column=COL_N_COMPROMISSOS).value,
            esperado.n_compromissos,
        )
        prazo = ws.cell(row=linha, column=COL_PRAZO_INICIAL)
        data = prazo.value.date() if isinstance(prazo.value, datetime) else prazo.value
        self.assertEqual(data, esperado.prazo_inicial)
        self.assertEqual(prazo.number_format, FORMATO_DATA)
        self.assertEqual(
            ws.cell(row=linha, column=COL_SITUACAO).value, esperado.situacao_atual
        )

    def test_xlsx_derivadas_de_prazo_tem_tipos_nativos(self):
        # `prazo_inicial`/`prazo_efetivo` são `date` nativo com formato
        # dd/mm/aaaa; `dias_atraso`/`n_prorrogacoes` são `int` nativo,
        # nunca texto (quebraria SUM/ordenação no Excel). A fixture de
        # exemplo não tem nenhum item real com remarcação pronta, então o
        # cenário (atrasado + 4 promessas distintas -> 3 remarcações) é
        # fabricado aqui sobre o item 17 (bloco ATRASADO, já tem
        # `prazo_entrega` vencido -> `dias_atraso` real).
        from apps.catalogo.models import SituacaoNormalizada
        from apps.pca.models import Acompanhamento, TipoEvento

        processo_17 = Processo.objects.get(item_pca=17)
        situacao_cat = SituacaoNormalizada.objects.first()
        for indice, prazo in enumerate(
            [date(2024, 2, 1), date(2024, 3, 1), date(2024, 4, 1), date(2024, 5, 1)]
        ):
            Acompanhamento.objects.create(
                processo=processo_17,
                referencia_data=date(2024, 1, 1) + timedelta(days=indice),
                origem_hash=f"hash-prazo-tipo-{processo_17.pk}-{indice}",
                tipo_evento=TipoEvento.REUNIAO_ACOMPANHAMENTO,
                situacao=situacao_cat,
                prazo_prometido=prazo,
            )

        esperado = Processo.objects.para_listagem().get(item_pca=17)
        _, ws = self._xlsx()
        linha = self._linha_xlsx_do_item(ws, 17)

        prazo_efetivo = ws.cell(row=linha, column=COL_PRAZO_EFETIVO)
        data = (
            prazo_efetivo.value.date()
            if isinstance(prazo_efetivo.value, datetime)
            else prazo_efetivo.value
        )
        self.assertEqual(data, esperado.prazo_efetivo)
        self.assertEqual(prazo_efetivo.number_format, FORMATO_DATA)

        dias_atraso = ws.cell(row=linha, column=COL_DIAS_ATRASO)
        self.assertIsInstance(dias_atraso.value, int)
        self.assertNotIsInstance(dias_atraso.value, bool)
        self.assertEqual(dias_atraso.value, esperado.dias_atraso)

        n_prorrogacoes = ws.cell(row=linha, column=COL_N_PRORROGACOES)
        self.assertIsInstance(n_prorrogacoes.value, int)
        self.assertNotIsInstance(n_prorrogacoes.value, bool)
        self.assertEqual(n_prorrogacoes.value, esperado.n_prorrogacoes)
        self.assertEqual(n_prorrogacoes.value, 3)

        # Um processo real sem promessa nenhuma (nem `prazo_entrega`, nem
        # histórico) exporta as duas colunas de prazo vazias (célula
        # real) e não o rótulo da tela — `prazo_inicial__isnull=True` já
        # cobre as duas fontes do COALESCE.
        sem_promessa = Processo.objects.para_listagem().filter(
            prazo_inicial__isnull=True
        ).first()
        if sem_promessa is not None:
            linha_sp = self._linha_xlsx_do_item(ws, sem_promessa.item_pca)
            self.assertIsNone(
                ws.cell(row=linha_sp, column=COL_PRAZO_INICIAL).value
            )
            self.assertIsNone(
                ws.cell(row=linha_sp, column=COL_PRAZO_EFETIVO).value
            )

    def test_xlsx_estado_situacao_batem_com_situacao_efetiva_ou_cancelado(self):
        # Mesma prova do CSV, sobre células XLSX: "Situação" reflete
        # `situacao_efetiva`, exceto quando `estado` é cancelado (aí é
        # sempre "Cancelado").
        processo = Processo.objects.get(item_pca=1)
        situacao_efetiva = Processo.objects.para_listagem().get(
            item_pca=1
        ).situacao_efetiva
        _, ws = self._xlsx()
        linha = self._linha_xlsx_do_item(ws, 1)
        self.assertEqual(
            ws.cell(row=linha, column=COL_ESTADO).value, processo.get_estado_display()
        )
        if processo.estado == Estado.CANCELADO.value:
            esperado_situacao = "Cancelado"
        else:
            esperado_situacao = Situacao(situacao_efetiva).label
        self.assertEqual(
            ws.cell(row=linha, column=COL_SITUACAO_PROCESSO).value,
            esperado_situacao,
        )

    def test_xlsx_neutraliza_injecao_de_formula(self):
        base = Processo.objects.first()
        Processo.objects.create(
            item_pca=901,
            descricao_objeto="=1+1",
            tipo=base.tipo,
            categoria=base.categoria,
            unidade_organizacional=base.unidade_organizacional,
            exercicio=base.exercicio,
        )
        _, ws = self._xlsx()
        linha = self._linha_xlsx_do_item(ws, 901)
        indice_objeto = CABECALHOS_ESPERADOS.index("Objeto") + 1
        celula = ws.cell(row=linha, column=indice_objeto)
        # Sem o apóstrofo, a célula viraria fórmula avaliável no pacote.
        self.assertIsInstance(celula.value, str)
        self.assertTrue(celula.value.startswith("'"))

    def test_xlsx_tem_mesmos_ids_da_tabela_com_querystring_combinado(self):
        params, esperado = self._filtro_combinado()
        ids_tabela = self._ids_tabela(**params)
        _, ws = self._xlsx(**params)
        ids_xlsx = [
            ws.cell(row=r, column=1).value
            for r in range(2, ws.max_row)  # exclui cabeçalho e linha de total
        ]
        self.assertEqual(len(ids_xlsx), esperado)
        self.assertEqual(ids_xlsx, ids_tabela)


class TestExportacaoOrdemCanonicaParidadeComATela(SimpleTestCase):
    """`_colunas_resolvidas` (exportação) e `colunas_selecionadas`
    (tela) precisam devolver a mesma ordem para a mesma seleção
    embaralhada — os dois reusam `apps.pca.colunas.canonizar`, então a
    paridade tabela/XLSX/CSV é estrutural, não coincidência de dado de
    teste. `SimpleTestCase`: sem banco, `_colunas_resolvidas`/
    `colunas_selecionadas` só leem o registro fechado e a
    requisição/sessão."""

    def test_ordem_do_export_bate_com_a_ordem_da_tela_para_selecao_embaralhada(self):
        chaves_embaralhadas = [
            "situacao",
            "modalidade",
            "item_pca",
            "descricao_objeto",
            "valor_estimado",
        ]
        request = RequestFactory().get(
            "/tabela", {"colunas": chaves_embaralhadas}
        )
        request.session = SessionStore()
        ordem_tela = colunas_selecionadas(request)

        ordem_export = [
            chave for chave, _, _ in _colunas_resolvidas(chaves_embaralhadas)
        ]

        self.assertEqual(ordem_tela, ordem_export)
        self.assertEqual(
            ordem_tela,
            [
                "item_pca",
                "descricao_objeto",
                "valor_estimado",
                "situacao",
                "modalidade",
            ],
        )


class TestExportacaoEspelhaTabela(TransactionTestCase):
    """O download espelha exatamente a filtragem, a ordem e as colunas
    que a tabela mostra no momento do clique."""

    serialized_rollback = True

    def setUp(self):
        call_command(
            "importar_pca",
            ARQUIVO_REAL,
            "--usuario=importador@pca.local",
            "--exercicio=2026",
            stdout=StringIO(),
        )
        self.usuario = get_user_model().objects.create_user(
            email="leitor@pca.local", password="x-forte-123"
        )
        self.client.force_login(self.usuario)

    def _ids_tabela_todas_paginas(self, **params):
        ids = []
        pagina_num = 1
        while True:
            resposta = self.client.get(
                reverse("pca:tabela"), {**params, "pagina": pagina_num}
            )
            pagina = resposta.context["pagina"]
            ids.extend(p.item_pca for p in pagina.object_list)
            if not pagina.has_next():
                break
            pagina_num += 1
        return ids

    def test_ordenacao_do_xlsx_bate_com_a_ordenacao_da_tabela(self):
        # `GET /tabela?ordenar=valor_estimado&dir=desc` e
        # `GET /exportar/xlsx?ordenar=valor_estimado&dir=desc` (mesma
        # querystring) produzem a MESMA ordem (achado 1 do plano —
        # `_queryset_export` não aplicava `_ordenar_para_tabela` antes).
        params = {"ordenar": "valor_estimado", "dir": "desc"}
        ids_tabela = self._ids_tabela_todas_paginas(**params)
        resposta = self.client.get(
            reverse("pca:exportar_xlsx"),
            {**params, "colunas": ["item_pca"]},
        )
        ws = load_workbook(BytesIO(resposta.content)).active
        # Só `item_pca` (TIPO_INTEIRO) selecionado: nenhuma coluna
        # monetária, logo nenhuma linha de total — `ws.max_row` é a
        # ÚLTIMA linha de dado (não uma linha de total a excluir).
        ids_xlsx = [
            ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)
        ]
        self.assertEqual(ids_xlsx, ids_tabela)

    def test_colunas_repetidas_restringem_e_ordenam_o_arquivo(self):
        resposta = self.client.get(
            reverse("pca:exportar_xlsx"),
            {"colunas": ["item_pca", "situacao"]},
        )
        ws = load_workbook(BytesIO(resposta.content)).active
        cabecalho = [ws.cell(row=1, column=c).value for c in (1, 2)]
        self.assertEqual(cabecalho, ["Item", "Situação"])
        self.assertIsNone(ws.cell(row=1, column=3).value)

    def test_sem_colunas_usa_o_padrao_compartilhado_da_fase(self):
        # Sem `?colunas=` e sessão vazia: 5 colunas padrão de
        # `apps.pca.colunas.COLUNAS_PADRAO`.
        resposta = self.client.get(reverse("pca:exportar_xlsx"))
        ws = load_workbook(BytesIO(resposta.content)).active
        cabecalho = [
            ws.cell(row=1, column=c).value for c in range(1, len(COLUNAS_PADRAO) + 1)
        ]
        mapa = {chave: cabecalho_ for chave, cabecalho_, _ in COLUNAS}
        self.assertEqual(cabecalho, [mapa[chave] for chave in COLUNAS_PADRAO])
        self.assertIsNone(
            ws.cell(row=1, column=len(COLUNAS_PADRAO) + 1).value
        )

    def test_link_exportar_xlsx_reflete_filtro_apos_troca_htmx(self):
        # Depois de um filtro por HTMX, o link "Exportar XLSX" presente no
        # fragmento devolvido contém o filtro atual na querystring.
        resposta = self.client.get(
            reverse("pca:tabela"),
            {"estado": Estado.CANCELADO.value},
            HTTP_HX_REQUEST="true",
        )
        conteudo = resposta.content.decode("utf-8")
        self.assertNotIn("hx-swap-oob", conteudo)
        hrefs = re.findall(
            rf'href="({re.escape(reverse("pca:exportar_xlsx"))}[^"]*)"', conteudo
        )
        self.assertEqual(len(hrefs), 1)
        parametros = dict(parse_qsl(urlparse(hrefs[0]).query))
        self.assertEqual(parametros.get("estado"), Estado.CANCELADO.value)

    def test_botao_exportar_xlsx_unico_e_com_icone_correto(self):
        resposta = self.client.get(reverse("pca:tabela"))
        conteudo = resposta.content.decode("utf-8")
        # A página intermediária (`pca:exportar`, sem sufixo `/csv`/`/xlsx`)
        # não é mais alcançada por nenhum link da tabela — checagem exata
        # (`?` logo após), já que "/exportar" também é PREFIXO de
        # "/exportar/xlsx" (falso positivo de simples substring).
        self.assertNotRegex(
            conteudo, rf'href="{re.escape(reverse("pca:exportar"))}\?'
        )
        hrefs = re.findall(
            rf'href="({re.escape(reverse("pca:exportar_xlsx"))}[^"]*)"', conteudo
        )
        self.assertEqual(len(hrefs), 1)
        self.assertIn("Exportar XLSX", conteudo)
        self.assertIn("fa-file-excel", conteudo)


class TestExportacaoNomeEAbaPorExercicio(TransactionTestCase):
    """Nome do arquivo e título da aba derivados do exercício efetivo,
    nunca "2026" fixo para outro exercício.

    Fixture leve via ORM direto: `importar_pca` exige um exercício
    previamente criado, então o exercício "2025" precisa existir antes
    via `Exercicio.objects.create` — o exercício "2026" já vem de uma
    migração de dados, preservado entre testes por
    `serialized_rollback`."""

    serialized_rollback = True

    def setUp(self):
        self.usuario = get_user_model().objects.create_user(
            email="leitor@pca.local", password="x-forte-123"
        )
        self.client.force_login(self.usuario)
        self.exercicio_2026 = Exercicio.objects.get(ano=2026)
        self.exercicio_2025 = Exercicio.objects.create(
            ano=2025, rotulo="PCA 2025", situacao=SituacaoExercicio.FECHADO
        )
        self.tipo = Tipo.objects.create(nome="Tipo 28-04")
        self.categoria = Categoria.objects.create(nome="Categoria 28-04")
        self.unidade = Unidade.objects.create(nome="Unidade 28-04")

    def _criar_processo(self, exercicio):
        Processo.objects.create(
            item_pca=1,
            descricao_objeto=f"Processo {exercicio.ano}",
            tipo=self.tipo,
            categoria=self.categoria,
            unidade_organizacional=self.unidade,
            exercicio=exercicio,
        )

    def test_nome_e_aba_seguem_o_exercicio_2025(self):
        self._criar_processo(self.exercicio_2025)
        resposta = self.client.get(
            reverse("pca:exportar_xlsx"), {"exercicio": "2025"}
        )
        self.assertIn("pca-2025-processos.xlsx", resposta["Content-Disposition"])
        ws = load_workbook(BytesIO(resposta.content)).active
        self.assertIn("2025", ws.title)
        self.assertNotIn("2026", ws.title)

    def test_nome_e_aba_seguem_o_exercicio_2026(self):
        self._criar_processo(self.exercicio_2026)
        resposta = self.client.get(
            reverse("pca:exportar_xlsx"), {"exercicio": "2026"}
        )
        self.assertIn("pca-2026-processos.xlsx", resposta["Content-Disposition"])
        ws = load_workbook(BytesIO(resposta.content)).active
        self.assertIn("2026", ws.title)

    def test_csv_nome_tambem_segue_o_exercicio(self):
        self._criar_processo(self.exercicio_2025)
        resposta = self.client.get(
            reverse("pca:exportar_csv"), {"exercicio": "2025"}
        )
        self.assertIn("pca-2025-processos.csv", resposta["Content-Disposition"])


class TestExportacaoZeroResultados(TransactionTestCase):
    """Filtro que não bate nada gera cabeçalho + indicação de ausência,
    nunca uma fórmula SUM sobre intervalo invertido."""

    serialized_rollback = True

    def setUp(self):
        call_command(
            "importar_pca",
            ARQUIVO_REAL,
            "--usuario=importador@pca.local",
            "--exercicio=2026",
            stdout=StringIO(),
        )
        self.usuario = get_user_model().objects.create_user(
            email="leitor@pca.local", password="x-forte-123"
        )
        self.client.force_login(self.usuario)

    def test_xlsx_sem_resultados_tem_cabecalho_mais_aviso_sem_total(self):
        resposta = self.client.get(
            reverse("pca:exportar_xlsx"),
            {"q": "termo-que-nao-bate-em-nenhum-processo-real-xyz123"},
        )
        ws = load_workbook(BytesIO(resposta.content)).active
        self.assertEqual(ws.max_row, 2)  # cabeçalho + aviso, sem total
        self.assertIsNotNone(ws.cell(row=2, column=1).value)
        self.assertNotEqual(ws.cell(row=2, column=1).value, "Total")
        # Nenhuma célula do arquivo tem fórmula (nenhum SUM invertido).
        for linha in ws.iter_rows():
            for celula in linha:
                if isinstance(celula.value, str):
                    self.assertFalse(celula.value.startswith("=SUM"))

    def test_csv_sem_resultados_tem_so_cabecalho(self):
        resposta = self.client.get(
            reverse("pca:exportar_csv"),
            {"q": "termo-que-nao-bate-em-nenhum-processo-real-xyz123"},
        )
        linhas = list(
            csv.reader(
                StringIO(resposta.content.decode("utf-8-sig")), delimiter=";"
            )
        )
        self.assertEqual(len(linhas), 1)


class TestExportacaoColunasSelecionadas(TransactionTestCase):
    """Registro único `apps.pca.colunas` — servidor: `?colunas=`
    repetido restringe e reordena o subconjunto de `exportacao.COLUNAS`
    exportado, com fallback para o padrão compartilhado quando o
    parâmetro está ausente."""

    serialized_rollback = True

    def setUp(self):
        call_command(
            "importar_pca",
            ARQUIVO_REAL,
            "--usuario=importador@pca.local",
            "--exercicio=2026",
            stdout=StringIO(),
        )
        self.usuario = get_user_model().objects.create_user(
            email="leitor@pca.local", password="x-forte-123"
        )
        self.client.force_login(self.usuario)

    def _cabecalhos_do_mapa(self, chaves):
        mapa = {chave: cabecalho for chave, cabecalho, _ in COLUNAS}
        return [mapa[chave] for chave in chaves]

    def test_sem_colunas_exporta_o_padrao_compartilhado_na_ordem_de_colunas(self):
        # O export (`pca:exportar_csv`/`pca:exportar_xlsx`) usa o mesmo
        # padrão compartilhado de `apps.pca.colunas.COLUNAS_PADRAO`, sem
        # conjunto próprio.
        esperado = self._cabecalhos_do_mapa(list(COLUNAS_PADRAO))
        self.assertEqual(
            esperado,
            ["Item", "Objeto", "Valor previsto (R$)", "Prazo atual", "Situação"],
        )

        resposta_csv = self.client.get(reverse("pca:exportar_csv"))
        linhas = list(
            csv.reader(
                StringIO(resposta_csv.content.decode("utf-8-sig")), delimiter=";"
            )
        )
        self.assertEqual(linhas[0], esperado)

        resposta_xlsx = self.client.get(reverse("pca:exportar_xlsx"))
        ws = load_workbook(BytesIO(resposta_xlsx.content)).active
        cabecalho = [ws.cell(row=1, column=c).value for c in range(1, len(esperado) + 1)]
        self.assertEqual(cabecalho, esperado)
        self.assertEqual(ws.max_row, 32)  # 30 processos + cabeçalho + total
        col_valor = esperado.index("Valor previsto (R$)") + 1
        letra_valor = get_column_letter(col_valor)
        celula_total = ws.cell(row=32, column=col_valor)
        self.assertEqual(
            celula_total.value, f"=SUM({letra_valor}2:{letra_valor}31)"
        )

    def test_ordem_recebida_e_canonizada_mesmo_fora_da_ordem_do_registro(self):
        # `estado`/`unidade_organizacional` são opcionais, `item_pca` é
        # núcleo — o resultado precisa trazer
        # `item_pca` primeiro (núcleo), depois as opcionais na ordem do
        # registro (`unidade_organizacional` antes de `estado`, Seção
        # Planejamento vem antes de Processo), independentemente da ordem
        # em que o `?colunas=` chegou.
        resposta = self.client.get(
            reverse("pca:exportar_csv"),
            {"colunas": ["estado", "item_pca", "unidade_organizacional"]},
        )
        linhas = list(
            csv.reader(StringIO(resposta.content.decode("utf-8-sig")), delimiter=";")
        )
        self.assertEqual(linhas[0], ["Item", "Unidade", "Estado"])

    def test_padrao_mais_extra_exporta_exatamente_seis_colunas(self):
        resposta = self.client.get(
            reverse("pca:exportar_csv"),
            {
                "colunas": [
                    "item_pca",
                    "descricao_objeto",
                    "unidade_organizacional",
                    "valor_estimado",
                    "estado",
                    "mes_previsto",
                ]
            },
        )
        linhas = list(
            csv.reader(StringIO(resposta.content.decode("utf-8-sig")), delimiter=";")
        )
        # Núcleo (item_pca, descricao_objeto, valor_estimado) primeiro,
        # na ordem de `COLUNAS_PADRAO`; depois as opcionais
        # (unidade_organizacional, mes_previsto, estado) na ordem do
        # registro `COLUNAS` (Planejamento antes de Processo).
        self.assertEqual(
            linhas[0],
            [
                "Item",
                "Objeto",
                "Valor previsto (R$)",
                "Unidade",
                "Mês previsto",
                "Estado",
            ],
        )

    def test_chave_invalida_e_descartada_em_silencio_sem_500(self):
        # `status` também é descartado em silêncio agora — não é mais um
        # allowlist válido de `exportacao.COLUNAS` (mesma consequência de
        # `campo_inventado`, uma chave que nunca existiu).
        resposta = self.client.get(
            reverse("pca:exportar_csv"),
            {"colunas": ["item_pca", "campo_inventado", "status", "estado"]},
        )
        self.assertEqual(resposta.status_code, 200)
        linhas = list(
            csv.reader(StringIO(resposta.content.decode("utf-8-sig")), delimiter=";")
        )
        self.assertEqual(linhas[0], ["Item", "Estado"])

    def test_valor_contratado_atrasado_e_publicacao_pendente_exportam(self):
        resposta = self.client.get(
            reverse("pca:exportar_csv"),
            {
                "colunas": [
                    "item_pca",
                    "valor_contratado",
                    "atrasado",
                    "publicacao_pendente",
                ]
            },
        )
        self.assertEqual(resposta.status_code, 200)
        linhas = list(
            csv.reader(StringIO(resposta.content.decode("utf-8-sig")), delimiter=";")
        )
        # Só `item_pca` é núcleo; as três opcionais entram na ordem do
        # registro `COLUNAS` (Processo: atrasado/publicacao_pendente,
        # antes de Execução contratual: valor_contratado).
        self.assertEqual(
            linhas[0],
            ["Item", "Atraso", "Publicação", "Valor contratado (R$)"],
        )

        processo_atrasado = (
            Processo.objects.para_listagem().filter(atrasado=True).first()
        )
        self.assertIsNotNone(processo_atrasado, "dado real tem ao menos 1 atrasado")
        linha_atrasado = self._linha_csv_do_item(linhas, processo_atrasado.item_pca)
        self.assertEqual(linha_atrasado[1], "Sim")

        processo_em_dia = (
            Processo.objects.para_listagem().filter(atrasado=False).first()
        )
        self.assertIsNotNone(processo_em_dia, "dado real tem ao menos 1 em dia")
        linha_em_dia = self._linha_csv_do_item(linhas, processo_em_dia.item_pca)
        self.assertEqual(linha_em_dia[1], "Não")

    def _linha_csv_do_item(self, linhas, item_pca):
        for linha in linhas[1:]:
            if linha[0] == str(item_pca):
                return linha
        self.fail(f"item {item_pca} não encontrado no CSV exportado")

    def test_modalidade_instrumento_e_publicacao_exportam(self):
        # As colunas novas de Execução contratual/Publicação exportam
        # pelo mesmo contrato das demais opcionais: FK pelo nome (mesma
        # regra de categoria/tipo), data solta como célula de data.
        resposta = self.client.get(
            reverse("pca:exportar_csv"),
            {
                "colunas": [
                    "item_pca",
                    "modalidade",
                    "instrumento_contratual",
                    "data_lancamento_spw",
                ]
            },
        )
        self.assertEqual(resposta.status_code, 200)
        linhas = list(
            csv.reader(StringIO(resposta.content.decode("utf-8-sig")), delimiter=";")
        )
        self.assertEqual(
            linhas[0],
            ["Item", "Modalidade", "Instrumento contratual", "Data lançamento SPW"],
        )

        processo_com_modalidade = (
            Processo.objects.para_listagem().filter(modalidade__isnull=False).first()
        )
        self.assertIsNotNone(
            processo_com_modalidade, "dado real tem ao menos 1 com modalidade"
        )
        linha = self._linha_csv_do_item(linhas, processo_com_modalidade.item_pca)
        self.assertEqual(linha[1], processo_com_modalidade.modalidade.nome)

        processo_com_instrumento = (
            Processo.objects.para_listagem()
            .filter(instrumento_contratual__isnull=False)
            .first()
        )
        self.assertIsNotNone(
            processo_com_instrumento, "dado real tem ao menos 1 com instrumento"
        )
        linha = self._linha_csv_do_item(linhas, processo_com_instrumento.item_pca)
        self.assertEqual(
            linha[2], processo_com_instrumento.instrumento_contratual.nome
        )

        processo_sem_modalidade = (
            Processo.objects.para_listagem().filter(modalidade__isnull=True).first()
        )
        self.assertIsNotNone(
            processo_sem_modalidade, "dado real tem ao menos 1 sem modalidade"
        )
        linha = self._linha_csv_do_item(linhas, processo_sem_modalidade.item_pca)
        self.assertEqual(linha[1], "")  # ausência é célula vazia, não rótulo

    def test_colunas_vazia_ou_desconhecida_cai_no_padrao_sem_500(self):
        # `?colunas=` sem sobrevivente válido (vazio ou só chave
        # desconhecida) não pode derrubar o endpoint nem produzir arquivo
        # sem coluna nenhuma: cai no mesmo padrão compartilhado do
        # caminho "sem `?colunas=`".
        esperado = self._cabecalhos_do_mapa(list(COLUNAS_PADRAO))
        for querystring in (
            {"colunas": [""]},
            {"colunas": ["xyz-invalido"]},
            {"colunas": ["xyz-invalido", ""]},
        ):
            with self.subTest(querystring=querystring):
                resposta_csv = self.client.get(
                    reverse("pca:exportar_csv"), querystring
                )
                self.assertEqual(resposta_csv.status_code, 200)
                linhas = list(
                    csv.reader(
                        StringIO(resposta_csv.content.decode("utf-8-sig")),
                        delimiter=";",
                    )
                )
                self.assertEqual(linhas[0], esperado)

                resposta_xlsx = self.client.get(
                    reverse("pca:exportar_xlsx"), querystring
                )
                self.assertEqual(resposta_xlsx.status_code, 200)
                ws = load_workbook(BytesIO(resposta_xlsx.content)).active
                cabecalho = [
                    ws.cell(row=1, column=c).value
                    for c in range(1, len(esperado) + 1)
                ]
                self.assertEqual(cabecalho, esperado)

    def test_colunas_repetidas_deduplicam_preservando_primeira_ocorrencia(self):
        # `?colunas=item_pca&colunas=item_pca` não pode duplicar a
        # coluna "Item" no arquivo exportado.
        resposta = self.client.get(
            reverse("pca:exportar_csv"),
            {"colunas": ["item_pca", "item_pca", "estado", "item_pca"]},
        )
        self.assertEqual(resposta.status_code, 200)
        linhas = list(
            csv.reader(StringIO(resposta.content.decode("utf-8-sig")), delimiter=";")
        )
        self.assertEqual(linhas[0], ["Item", "Estado"])

    def test_xlsx_sem_coluna_monetaria_nao_ganha_linha_de_total(self):
        resposta = self.client.get(
            reverse("pca:exportar_xlsx"), {"colunas": ["item_pca", "estado"]}
        )
        ws = load_workbook(BytesIO(resposta.content)).active
        self.assertEqual(ws.max_row, 31)  # 30 processos + cabeçalho, sem total

    def test_xlsx_colunas_embaralhadas_cabecalho_comeca_por_item_e_objeto(self):
        # `?colunas=` real de ponta a ponta com o núcleo fora de ordem e
        # misturado a opcionais precisa produzir um XLSX cujo cabeçalho
        # começa pelo núcleo padrão (Item, Objeto), nunca pela ordem
        # crua recebida.
        resposta = self.client.get(
            reverse("pca:exportar_xlsx"),
            {
                "colunas": [
                    "situacao",
                    "modalidade",
                    "item_pca",
                    "descricao_objeto",
                    "valor_estimado",
                ]
            },
        )
        self.assertEqual(resposta.status_code, 200)
        ws = load_workbook(BytesIO(resposta.content)).active
        cabecalho = [ws.cell(row=1, column=c).value for c in range(1, 6)]
        self.assertEqual(
            cabecalho,
            [
                "Item",
                "Objeto",
                "Valor previsto (R$)",
                "Situação",
                "Modalidade",
            ],
        )


class TestExportacaoFormularioSoXlsx(TransactionTestCase):
    """A interface de Processos oferece somente XLSX; `pca/exportar.html`
    não apresenta CSV como opção visível (a rota `pca:exportar_csv`
    continua respondendo, mas nunca a partir de um botão/formulário)."""

    serialized_rollback = True

    def setUp(self):
        call_command(
            "importar_pca",
            ARQUIVO_REAL,
            "--usuario=importador@pca.local",
            "--exercicio=2026",
            stdout=StringIO(),
        )
        self.usuario = get_user_model().objects.create_user(
            email="leitor@pca.local", password="x-forte-123"
        )
        self.client.force_login(self.usuario)

    def test_pagina_exportar_nao_mostra_csv_como_opcao_visivel(self):
        resposta = self.client.get(reverse("pca:exportar"))
        conteudo = resposta.content.decode("utf-8")
        self.assertNotIn(">CSV<", conteudo)
        self.assertNotIn("Formato", conteudo)
        # O campo continua presente, só oculto (compatibilidade do POST).
        self.assertIn('type="hidden" name="formato" value="xlsx"', conteudo)

    def test_post_na_pagina_exportar_com_o_campo_oculto_gera_xlsx(self):
        # O navegador sempre envia o campo `hidden` junto do POST (é o que
        # `{{ form }}` já renderiza, valor "xlsx" fixo) — o teste simula
        # exatamente esse POST real, não a ausência do campo (que apenas
        # invalidaria o formulário por ser `required`, cenário à parte).
        resposta = self.client.post(
            reverse("pca:exportar"),
            {"colunas": ["item_pca"], "formato": "xlsx"},
        )
        self.assertEqual(
            resposta["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )


class TestExportacaoBotoes(TransactionTestCase):
    """O link "Exportar XLSX" na tabela: download direto, sem a página
    intermediária de formato/colunas, sempre carregando o querystring
    validado (querystring_filtros, nunca o GET cru) e atualizado a cada
    troca de filtro HTMX, senão o botão exportaria um filtro que já não é
    o da tela."""

    serialized_rollback = True

    def setUp(self):
        call_command(
            "importar_pca",
            ARQUIVO_REAL,
            "--usuario=importador@pca.local",
            "--exercicio=2026",
            stdout=StringIO(),
        )
        self.usuario = get_user_model().objects.create_user(
            email="leitor@pca.local", password="x-forte-123"
        )
        self.client.force_login(self.usuario)

    def test_tabela_tem_botao_exportar_xlsx_unico_apontando_direto_ao_download(self):
        resposta = self.client.get(reverse("pca:tabela"))
        conteudo = resposta.content.decode("utf-8")
        self.assertNotIn('id="acoes-export"', conteudo)
        hrefs = re.findall(
            rf'href="({re.escape(reverse("pca:exportar_xlsx"))}[^"]*)"', conteudo
        )
        self.assertEqual(len(hrefs), 1)

    def test_botao_exportar_preserva_querystring_validado_e_ignora_o_cru(self):
        # `estado` é o eixo que `querystring_filtros` reconhece e
        # preserva hoje.
        resposta = self.client.get(
            reverse("pca:tabela"),
            {"estado": Estado.CANCELADO.value, "campo_inventado": "1; DROP TABLE"},
        )
        conteudo = resposta.content.decode("utf-8")
        hrefs = re.findall(
            rf'href="({re.escape(reverse("pca:exportar_xlsx"))}[^"]*)"', conteudo
        )
        self.assertEqual(len(hrefs), 1)
        parametros = dict(parse_qsl(urlparse(hrefs[0]).query))
        self.assertEqual(parametros.get("estado"), Estado.CANCELADO.value)
        self.assertNotIn("campo_inventado", parametros)

    def test_botao_exportar_atualiza_sem_oob_na_troca_de_filtro_htmx(self):
        """O botão "Exportar XLSX" mora dentro de `_tabela_resultado.html`
        (alvo do próprio `hx-get`), não no cabeçalho estático — por isso
        a querystring corrente chega sem nenhum `hx-swap-oob`."""
        resposta = self.client.get(
            reverse("pca:tabela"),
            {"estado": Estado.CANCELADO.value},
            HTTP_HX_REQUEST="true",
        )
        conteudo = resposta.content.decode("utf-8")
        self.assertNotIn("hx-swap-oob", conteudo)
        hrefs = re.findall(
            rf'href="({re.escape(reverse("pca:exportar_xlsx"))}[^"]*)"', conteudo
        )
        self.assertEqual(len(hrefs), 1)
        parametros = dict(parse_qsl(urlparse(hrefs[0]).query))
        self.assertEqual(parametros.get("estado"), Estado.CANCELADO.value)
