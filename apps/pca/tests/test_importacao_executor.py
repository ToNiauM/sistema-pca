"""Confirma que `executar_import`/`USUARIO_PADRAO` são importáveis do
módulo e que `EventoImportacao` está pronto (schema + integridade). A
prova comportamental do import continua em `test_import.py`.

`TestCabecalhoPrazoInicialAceito` prova que a coluna de origem do prazo
importa igual com os dois cabeçalhos aceitos ("Data prevista para
entrega", o histórico do modelo, e "Prazo inicial", o rótulo novo de
tela/XLSX): o executor lê só por posição, nunca por texto de cabeçalho."""

from datetime import date
from io import StringIO
from pathlib import Path
import tempfile

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import SimpleTestCase, TransactionTestCase
from openpyxl import Workbook

from apps.catalogo.models import Exercicio
from apps.pca.importacao.executor import COLUNAS_CATALOGO, USUARIO_PADRAO, executar_import
from apps.pca.importacao.modelo import (
    CABECALHOS_ACOMPANHAMENTO,
    CABECALHOS_AUXILIARES,
    CABECALHOS_PLANILHA1,
)
from apps.pca.importacao.models import EventoImportacao
from apps.pca.models import Processo

ARQUIVO_REAL = "apps/pca/fixtures/modelo-controle-exemplo.xlsx"


class TestImportavel(SimpleTestCase):
    def test_import_de_executar_import_e_usuario_padrao(self):
        self.assertTrue(callable(executar_import))
        self.assertEqual(USUARIO_PADRAO, "importador@pca.local")


class TestEventoImportacaoCampos(SimpleTestCase):
    def test_campos_esperados(self):
        nomes = [f.name for f in EventoImportacao._meta.get_fields()]
        for campo in [
            "disparado_por",
            "criado_em",
            "arquivo_nome",
            "exercicio",
            "processos_criados",
            "acompanhamentos_criados",
            "relatorio_caminho",
        ]:
            self.assertIn(campo, nomes)


class TestEventoImportacaoPersistencia(TransactionTestCase):
    serialized_rollback = True

    def test_criar_evento_importacao_nao_levanta_erro_de_integridade(self):
        usuario = get_user_model().objects.get(email=USUARIO_PADRAO)
        exercicio = Exercicio.objects.get(ano=2026)

        evento = EventoImportacao.objects.create(
            disparado_por=usuario,
            arquivo_nome="apps/pca/fixtures/modelo-controle-exemplo.xlsx",
            exercicio=exercicio,
            processos_criados=149,
            acompanhamentos_criados=381,
            relatorio_caminho="ops/relatorios/import-2026-08-05-2200.md",
        )

        self.assertIsNotNone(evento.pk)
        self.assertIsNotNone(evento.criado_em)


class TestExecutarImportCaminhoRelatorio(TransactionTestCase):
    """Mesmo padrão de `test_import.py`: `TransactionTestCase` porque
    `executar_import` gerencia sua própria transação, e `serialized_rollback`
    porque o usuário de serviço só existe via migração de dados."""

    serialized_rollback = True

    def test_relatorio_devolvido_tem_caminho_relatorio_setado(self):
        usuario = get_user_model().objects.get(email=USUARIO_PADRAO)
        exercicio = Exercicio.objects.get(ano=2026)

        relatorio = executar_import(
            ARQUIVO_REAL,
            usuario,
            exercicio,
            dry_run=False,
            escrever_stdout=StringIO().write,
        )

        self.assertTrue(hasattr(relatorio, "caminho_relatorio"))
        self.assertTrue(Path(relatorio.caminho_relatorio).exists())


class TestCabecalhoPrazoInicialAceito(TransactionTestCase):
    """`_importar_processos` lê a coluna de origem do prazo (posição 14
    de `Planilha1`) por índice, nunca por texto de cabeçalho: um arquivo
    com "Prazo inicial" na linha 1 importa `prazo_entrega` igual a um
    arquivo com o cabeçalho histórico "Data prevista para entrega".
    `CABECALHOS_PLANILHA1` (o modelo de download) continua emitindo só o
    cabeçalho histórico."""

    serialized_rollback = True

    def _construir_workbook(self, item_pca, cabecalho_coluna_14, prazo):
        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Planilha1"
        cabecalhos = list(CABECALHOS_PLANILHA1)
        cabecalhos[13] = cabecalho_coluna_14
        for coluna, cabecalho in enumerate(cabecalhos, start=1):
            planilha.cell(row=1, column=coluna, value=cabecalho)

        linha = [None] * 40
        linha[0] = item_pca
        linha[2] = f"Objeto sintético {item_pca}"
        linha[3] = "Justificativa sintética"
        linha[4] = "Nova Contratação"
        linha[5] = "Serviços"
        linha[6] = "Unidade Sintética Prazo Inicial"
        linha[7] = 1000
        linha[13] = prazo
        linha[16] = "Não iniciado"
        for coluna, valor in enumerate(linha, start=1):
            planilha.cell(row=2, column=coluna, value=valor)

        acompanhamento = workbook.create_sheet("Acompanhamento")
        for coluna, cabecalho in enumerate(CABECALHOS_ACOMPANHAMENTO, start=1):
            acompanhamento.cell(row=1, column=coluna, value=cabecalho)

        # `_upsert_catalogo` (executor.py) faz `cabecalho.index(nome_coluna)`
        # para TODAS as 8 entradas de `COLUNAS_CATALOGO` — o cabeçalho da
        # aba Listas precisa das 8 colunas presentes, mesmo com dado só
        # nas 3 primeiras (Tipo/Categoria/UO, as únicas que a linha
        # sintética referencia).
        listas = workbook.create_sheet("Listas")
        cabecalhos_listas = tuple(nome for nome, _ in COLUNAS_CATALOGO) + CABECALHOS_AUXILIARES
        for coluna, cabecalho in enumerate(cabecalhos_listas, start=1):
            listas.cell(row=1, column=coluna, value=cabecalho)
        listas.cell(row=2, column=1, value="Nova Contratação")  # Tipo de Contratação
        listas.cell(row=2, column=2, value="Serviços")  # Categoria
        listas.cell(row=2, column=3, value="Unidade Sintética Prazo Inicial")  # UO

        arquivo = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        workbook.save(arquivo.name)
        workbook.close()
        return Path(arquivo.name)

    def _importar(self, item_pca, cabecalho_coluna_14, prazo):
        # `_upsert_catalogo` cria Tipo/Categoria/Unidade a partir da aba
        # Listas — não é preciso pré-criar nada aqui.
        caminho = self._construir_workbook(item_pca, cabecalho_coluna_14, prazo)
        try:
            usuario = get_user_model().objects.get(email=USUARIO_PADRAO)
            exercicio = Exercicio.objects.get(ano=2026)
            call_command(
                "importar_pca", str(caminho), exercicio=exercicio.ano,
                usuario=usuario.email, stdout=StringIO(),
            )
        finally:
            caminho.unlink(missing_ok=True)

    def test_cabecalho_historico_e_cabecalho_novo_importam_o_mesmo_prazo(self):
        prazo = date(2026, 8, 1)
        self._importar(9101, "Data prevista para entrega", prazo)
        self._importar(9102, "Prazo inicial", prazo)

        exercicio = Exercicio.objects.get(ano=2026)
        historico = Processo.objects.get(exercicio=exercicio, item_pca=9101)
        novo = Processo.objects.get(exercicio=exercicio, item_pca=9102)
        self.assertEqual(historico.prazo_entrega, prazo)
        self.assertEqual(novo.prazo_entrega, prazo)
        self.assertEqual(historico.prazo_entrega, novo.prazo_entrega)

    def test_modelo_de_download_continua_emitindo_o_cabecalho_historico(self):
        # A planilha modelo não muda: só a leitura aceita o rótulo novo,
        # a emissão continua histórica.
        self.assertEqual(
            CABECALHOS_PLANILHA1[13], "Data prevista para entrega"
        )
