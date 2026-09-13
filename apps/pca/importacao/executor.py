"""Núcleo do import: a regra de import existe num único lugar reutilizável,
chamado tanto pelo comando de CLI quanto pela tela de importação do admin.

Nunca django-import-export com use_bulk=True, que desliga post_save/pre_save
e mata o histórico. Lê o .xlsx informado, normaliza via
apps.pca.importacao.parsers, grava com histórico sob o usuário de serviço,
é idempotente em reexecução e emite o relatório de conferência nominal
mesmo sob --dry-run."""

from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.management.base import CommandError
from django.db import transaction
from django.utils.timezone import localdate
from simple_history.utils import bulk_create_with_history

from apps.catalogo.models import (
    Categoria,
    Classificacao,
    GrauPrioridade,
    InstrumentoContratual,
    Modalidade,
    SituacaoNormalizada,
    Tipo,
    Unidade,
)
from apps.catalogo.utils import normalizar_nome
from apps.pca.importacao.parsers import (
    parse_data_excel,
    parse_mes_previsto,
    parse_numeros_sei,
    parse_status,
)
from apps.pca.importacao.relatorio import Relatorio
from apps.pca.models import (
    Acompanhamento,
    Estado,
    Processo,
    ProcessoSEI,
    Reuniao,
    Situacao,
    SituacaoReuniao,
)
from apps.pca.regras_situacao import situacao_inicial
from apps.pca.utils import calcular_origem_hash

# os mesmos 5 rótulos de STATUS_POR_ROTULO (parsers.py); o slug legado é
# usado aqui só para decidir estado, nunca gravado no banco
STATUS_LEGADO_VALIDOS = {
    "concluido", "nao_iniciado", "em_tramitacao", "vigente", "cancelado",
}

USUARIO_PADRAO = "importador@pca.local"
MOTIVO_HISTORICO = "import inicial"

# rótulo da coluna na aba Listas -> modelo de catálogo. Status e Situação
# do SEI são TextChoices no código, ficam fora deste registro
COLUNAS_CATALOGO = [
    ("Tipo de Contratação", Tipo),
    ("Categoria", Categoria),
    ("UO (Área Demandante)", Unidade),
    ("Grau de Prioridade", GrauPrioridade),
    ("Classificação", Classificacao),
    ("Modalidade", Modalidade),
    ("Instrumento Contratual", InstrumentoContratual),
    ("Situação Normalizada", SituacaoNormalizada),
]

# Situação do SEI e Tipo de Evento são TextChoices de apps.pca.models,
# mapeados aqui por rótulo literal para não forçar parsers.py a importar
# django.db.models
SITUACAO_SEI_POR_ROTULO = {
    "Autuado": "autuado",
    "A autuar": "a_autuar",
    "Não se aplica": "nao_se_aplica",
}
TIPO_EVENTO_POR_ROTULO = {
    "Reunião de acompanhamento": "reuniao_acompanhamento",
    "Gestão de Riscos": "gestao_riscos",
    # acompanhamentos com tipo_evento="automatico" (gravados por
    # aplicar_regras_automaticas/correções retroativas) precisam do rótulo
    # "Automático" reconhecido de volta no reimport
    "Automático": "automatico",
}


def _texto(valor):
    """CharField tolerante a lixo de planilha; nunca deixa passar um tipo
    não string para o ORM, mas também nunca inventa conteúdo."""
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def _decimal(valor):
    if valor is None:
        return None
    try:
        return Decimal(str(valor))
    except InvalidOperation:
        return None


class ImportadorPca:
    """Orquestração do import para um exercicio específico."""

    def __init__(self, exercicio):
        self.exercicio = exercicio

    # -- Orquestração ------------------------------------------------------

    def importar(self, workbook, usuario, relatorio):
        mapas_catalogo = self._upsert_catalogo(workbook["Listas"])
        # a promessa vigente por item vem da planilha, não do banco:
        # Acompanhamento ainda não existe para os itens novos neste ponto
        mapa_prazo_vigente = self._mapa_prazo_vigente(workbook["Acompanhamento"])
        mapa_item_para_pk = self._importar_processos(
            workbook["Planilha1"], usuario, mapas_catalogo, relatorio,
            mapa_prazo_vigente,
        )
        self._importar_acompanhamentos(
            workbook["Acompanhamento"],
            usuario,
            mapas_catalogo,
            mapa_item_para_pk,
            relatorio,
        )

    def _mapa_prazo_vigente(self, aba_acompanhamento):
        """Varre a aba Acompanhamento uma vez e devolve {item_pca:
        prazo_prometido} só com a promessa de maior referencia_data por
        item. Linhas sem referencia_data ou prazo_prometido são ignoradas."""
        linhas = list(aba_acompanhamento.iter_rows(min_row=2, values_only=True))
        mapa_data_mais_recente = {}
        mapa_prazo = {}
        for linha in linhas:
            item_processo = linha[1]
            if item_processo is None:
                continue
            referencia_data = parse_data_excel(linha[3])
            prazo_prometido = parse_data_excel(linha[8])
            if referencia_data is None or prazo_prometido is None:
                continue
            data_atual = mapa_data_mais_recente.get(item_processo)
            if data_atual is None or referencia_data >= data_atual:
                mapa_data_mais_recente[item_processo] = referencia_data
                mapa_prazo[item_processo] = prazo_prometido
        return mapa_prazo

    # -- Catálogo: upsert por nome_normalizado, nunca apaga --------------

    def _upsert_catalogo(self, aba_listas):
        cabecalho = [c.value for c in next(aba_listas.iter_rows(min_row=1, max_row=1))]
        linhas = list(aba_listas.iter_rows(min_row=2, values_only=True))

        mapas = {}
        for nome_coluna, modelo in COLUNAS_CATALOGO:
            idx = cabecalho.index(nome_coluna)
            mapa_normalizado = {}
            for linha in linhas:
                valor = linha[idx] if idx < len(linha) else None
                if not valor:
                    continue
                chave = normalizar_nome(valor)
                if chave in mapa_normalizado:
                    continue
                instancia, criado = modelo.objects.get_or_create(
                    nome_normalizado=chave, defaults={"nome": valor}
                )
                if not criado and instancia.nome != valor:
                    instancia.nome = valor
                    instancia.save()
                mapa_normalizado[chave] = instancia
            mapas[modelo] = mapa_normalizado
        return mapas

    def _resolver_fk(self, mapa, valor, *, obrigatorio=False, item_pca=None, campo=None):
        if not valor:
            if obrigatorio:
                raise CommandError(
                    f"Processo {item_pca}: campo obrigatório '{campo}' vazio na planilha."
                )
            return None
        instancia = mapa.get(normalizar_nome(valor))
        if instancia is None:
            if obrigatorio:
                raise CommandError(
                    f"Processo {item_pca}: valor '{valor}' de '{campo}' não "
                    "consta na aba Listas."
                )
            return None
        return instancia

    # -- Processo: upsert por item_pca, só insere o que falta ------------

    def _importar_processos(
        self, aba_planilha1, usuario, mapas_catalogo, relatorio,
        mapa_prazo_vigente=None,
    ):
        # min_row=2 pula o cabeçalho por completo: nunca lido por texto,
        # só por posição fixa (linha[N], ver abaixo)
        mapa_prazo_vigente = mapa_prazo_vigente or {}
        linhas = list(aba_planilha1.iter_rows(min_row=2, values_only=True))
        existentes = dict(
            Processo.objects.filter(exercicio=self.exercicio).values_list(
                "item_pca", "pk"
            )
        )

        novos = []
        numeros_sei_por_item = {}

        for linha in linhas:
            item_pca = linha[0]
            if item_pca is None:
                continue  # linha em branco no rodapé da aba

            if item_pca in existentes:
                relatorio.registrar_processo_ignorado()
                continue

            mes_bruto = linha[8]
            mes_previsto = parse_mes_previsto(mes_bruto)
            if mes_bruto and mes_previsto is None:
                relatorio.registrar_mes_ambiguo(item_pca)

            data_envio = parse_data_excel(linha[12])
            data_entrega = parse_data_excel(linha[14])
            if data_envio and data_entrega and data_entrega < data_envio:
                relatorio.registrar_entrega_anterior_ao_envio(
                    item_pca, data_envio, data_entrega
                )

            grau_prioridade = self._resolver_fk(mapas_catalogo[GrauPrioridade], linha[10])
            classificacao = self._resolver_fk(mapas_catalogo[Classificacao], linha[11])
            if grau_prioridade is None and classificacao is None:
                relatorio.registrar_sem_prioridade_e_classificacao(item_pca)

            numeros_sei = parse_numeros_sei(linha[1])
            if len(numeros_sei) > 1:
                relatorio.registrar_multiplos_sei(item_pca, numeros_sei)
            numeros_sei_por_item[item_pca] = numeros_sei

            status_legado = parse_status(linha[16])
            if status_legado not in STATUS_LEGADO_VALIDOS:
                raise CommandError(
                    f"Processo {item_pca}: status '{linha[16]}' não é mais "
                    "suportado pelo sistema. Revise a linha na planilha "
                    "antes de importar."
                )

            tipo_resolvido = self._resolver_fk(
                mapas_catalogo[Tipo], linha[4],
                obrigatorio=True, item_pca=item_pca, campo="tipo de contratação",
            )
            estado_valor = (
                Estado.CANCELADO if status_legado == "cancelado" else Estado.ATIVO
            )
            situacao_valor = situacao_inicial(
                tipo_nome_normalizado=tipo_resolvido.nome_normalizado,
                data_assinatura_contrato=parse_data_excel(linha[35]),
                data_recebimento_gelic=data_entrega,
                prazo_entrega=parse_data_excel(linha[13]),
                prazo_vigente=mapa_prazo_vigente.get(item_pca),
                hoje=localdate(),
            )

            novos.append(
                Processo(
                    item_pca=item_pca,
                    exercicio=self.exercicio,
                    descricao_objeto=_texto(linha[2]) or "",
                    justificativa=_texto(linha[3]) or "",
                    tipo=tipo_resolvido,
                    categoria=self._resolver_fk(
                        mapas_catalogo[Categoria], linha[5],
                        obrigatorio=True, item_pca=item_pca, campo="categoria",
                    ),
                    unidade_organizacional=self._resolver_fk(
                        mapas_catalogo[Unidade], linha[6],
                        obrigatorio=True, item_pca=item_pca, campo="UO",
                    ),
                    valor_estimado=_decimal(linha[7]),
                    mes_previsto=mes_previsto,
                    data_inclusao_pca=parse_data_excel(linha[9]),
                    grau_prioridade=grau_prioridade,
                    classificacao=classificacao,
                    data_envio_gelic=data_envio,
                    prazo_entrega=parse_data_excel(linha[13]),
                    data_recebimento_gelic=data_entrega,
                    data_prevista_conclusao=parse_data_excel(linha[15]),
                    estado=estado_valor,
                    situacao=situacao_valor,
                    # bloco de execução contratual, colunas 25-39; a coluna
                    # 24 ("Data Recebimento do Processo") fundiu-se em
                    # data_recebimento_gelic acima e deixa de ser lida aqui
                    modalidade=self._resolver_fk(mapas_catalogo[Modalidade], linha[25]),
                    vigencia_inicio=parse_data_excel(linha[26]),
                    vigencia_fim=parse_data_excel(linha[27]),
                    numero_contratacao=_texto(linha[28]),
                    numero_arp=_texto(linha[29]),
                    instrumento_contratual=self._resolver_fk(
                        mapas_catalogo[InstrumentoContratual], linha[30]
                    ),
                    numero_instrumento_contratual=_texto(linha[31]),
                    valor_contratado=_decimal(linha[32]),
                    fornecedor_cnpj=_texto(linha[33]),
                    fornecedor_razao_social=_texto(linha[34]),
                    data_assinatura_contrato=parse_data_excel(linha[35]),
                    data_lancamento_spw=parse_data_excel(linha[36]),
                    data_lancamento_wordpress=parse_data_excel(linha[37]),
                    data_lancamento_dados_abertos=parse_data_excel(linha[38]),
                    situacao_sei=SITUACAO_SEI_POR_ROTULO.get(_texto(linha[39]) or "", ""),
                )
            )

        criados = bulk_create_with_history(
            novos,
            Processo,
            batch_size=200,
            default_user=usuario,
            default_change_reason=MOTIVO_HISTORICO,
        )
        # Postgres devolve PK preenchido via RETURNING, sem 2ª query
        mapa_item_para_pk = {
            (self.exercicio.pk, item_pca): pk for item_pca, pk in existentes.items()
        }
        mapa_item_para_pk.update(
            {(self.exercicio.pk, p.item_pca): p.pk for p in criados}
        )

        sei_novos = [
            ProcessoSEI(processo_id=processo.pk, numero_sei=numero)
            for processo in criados
            for numero in numeros_sei_por_item.get(processo.item_pca, [])
        ]
        if sei_novos:
            # ProcessoSEI não tem HistoricalRecords; bulk_create comum é correto aqui
            ProcessoSEI.objects.bulk_create(sei_novos)

        relatorio.registrar_processos_criados(len(criados))
        return mapa_item_para_pk

    # -- Acompanhamento: upsert por origem_hash ---------------------------

    def _reuniao_de(self, referencia_data, cache):
        """Uma Reuniao por data de referência distinta, criada sob demanda
        e memorizada na chamada. Fechada de propósito: são reuniões que já
        aconteceram."""
        if referencia_data not in cache:
            cache[referencia_data], _ = Reuniao.objects.get_or_create(
                data=referencia_data,
                defaults={
                    "situacao": SituacaoReuniao.FECHADA,
                    "exercicio": self.exercicio,
                },
            )
        return cache[referencia_data]

    def _importar_acompanhamentos(
        self, aba_acompanhamento, usuario, mapas_catalogo, mapa_item_para_pk, relatorio
    ):
        linhas = list(aba_acompanhamento.iter_rows(min_row=2, values_only=True))

        candidatos = []
        for linha in linhas:
            item_processo = linha[1]
            chave_processo = (self.exercicio.pk, item_processo)
            if item_processo is None or chave_processo not in mapa_item_para_pk:
                continue  # linha em branco, ou órfã de um processo inexistente
            referencia_data = parse_data_excel(linha[3])
            if referencia_data is None:
                continue  # defensivo: os registros reais sempre têm data
            situacao_informada = _texto(linha[6]) or ""
            origem_hash = calcular_origem_hash(
                self.exercicio.ano, item_processo, referencia_data, situacao_informada
            )
            candidatos.append(
                (
                    origem_hash,
                    chave_processo,
                    referencia_data,
                    situacao_informada,
                    linha,
                )
            )

        hashes_candidatos = [c[0] for c in candidatos]
        hashes_existentes = set(
            Acompanhamento.objects.filter(
                origem_hash__in=hashes_candidatos
            ).values_list("origem_hash", flat=True)
        )

        novos = []
        ignorados = 0
        # acompanhamentos criados pela UI usam um hash salgado com UUID, sem
        # relação com a chave natural; sem este set, duas linhas novas do
        # mesmo lote com o mesmo origem_hash recomputado estourariam
        # UniqueViolation no bulk_create. A idempotência por hash se estende
        # ao próprio lote: a segunda ocorrência é tratada como "já existe"
        hashes_no_lote = set()
        # associa a reunião aqui também para o reimport não deixar os
        # acompanhamentos avulsos e n_reunioes zerado na tabela/export
        reunioes = {}
        for origem_hash, chave_processo, referencia_data, situacao_informada, linha in candidatos:
            if origem_hash in hashes_existentes or origem_hash in hashes_no_lote:
                ignorados += 1
                continue
            hashes_no_lote.add(origem_hash)

            situacao = self._resolver_fk(
                mapas_catalogo[SituacaoNormalizada], linha[7],
                obrigatorio=True,
                item_pca=chave_processo[1],
                campo="situação normalizada",
            )
            novos.append(
                Acompanhamento(
                    processo_id=mapa_item_para_pk[chave_processo],
                    referencia_data=referencia_data,
                    reuniao=self._reuniao_de(referencia_data, reunioes),
                    origem_hash=origem_hash,
                    evento=_texto(linha[4]) or "",
                    tipo_evento=TIPO_EVENTO_POR_ROTULO.get(_texto(linha[5]) or "", ""),
                    situacao_informada=situacao_informada,
                    situacao=situacao,
                    prazo_prometido=parse_data_excel(linha[8]),
                    data_evento_informada=parse_data_excel(linha[9]),
                    area_informada=_texto(linha[10]) or "",
                )
            )

        bulk_create_with_history(
            novos,
            Acompanhamento,
            batch_size=500,
            default_user=usuario,
            default_change_reason=MOTIVO_HISTORICO,
        )
        relatorio.registrar_acompanhamentos_criados(len(novos))
        relatorio.registrar_acompanhamentos_ignorados(ignorados)


def executar_import(caminho, usuario, exercicio, dry_run, escrever_stdout=None):
    """Abre caminho, roda a importação inteira dentro de uma transação
    (com set_rollback(True) se dry_run) e devolve o Relatorio já
    preenchido. Levanta CommandError se o arquivo não abrir.

    relatorio.caminho_relatorio fica setado com o Path que
    Relatorio.emitir() devolveu."""
    try:
        workbook = openpyxl.load_workbook(caminho, data_only=True)
    except (FileNotFoundError, OSError) as exc:
        raise CommandError(f"Não foi possível abrir '{caminho}': {exc}") from exc

    relatorio = Relatorio(arquivo_fonte=caminho, dry_run=dry_run)
    importador = ImportadorPca(exercicio)

    with transaction.atomic():
        importador.importar(workbook, usuario, relatorio)
        caminho_relatorio = relatorio.emitir(escrever_stdout or (lambda texto: None))
        if dry_run:
            # preserva o relatório já montado/gravado; não persiste nada
            transaction.set_rollback(True)

    relatorio.caminho_relatorio = caminho_relatorio
    return relatorio
